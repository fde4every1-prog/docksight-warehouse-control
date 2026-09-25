"""Generate the DockSight Improvised Version stakeholder presentation."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = Path(__file__).resolve().parent / "DockSight_Improvised_Version.pptx"

# DockSight V3: midnight control-room, steel, signal amber, safe green.
BG = RGBColor(0x0D, 0x17, 0x20)
PANEL = RGBColor(0x15, 0x24, 0x30)
PANEL_ALT = RGBColor(0x19, 0x2E, 0x3B)
LINE = RGBColor(0x2E, 0x47, 0x55)
TEXT = RGBColor(0xF1, 0xF5, 0xF7)
MUTED = RGBColor(0x9E, 0xAF, 0xB9)
CYAN = RGBColor(0x42, 0xA6, 0xC6)
AMBER = RGBColor(0xD4, 0x9A, 0x32)
GREEN = RGBColor(0x4E, 0xA7, 0x78)
RED = RGBColor(0xD1, 0x62, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_run(run, text, size=16, bold=False, color=TEXT):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Aptos"


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def send_back(slide, shape):
    tree = slide.shapes._spTree
    element = shape._element
    tree.remove(element)
    tree.insert(2, element)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def chrome(slide, number, section):
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5)
    )
    fill(bg, BG)
    send_back(slide, bg)

    rail = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), Inches(7.5)
    )
    fill(rail, CYAN)

    footer = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Inches(7.16), Inches(13.333), Inches(0.34)
    )
    fill(footer, PANEL)

    tb = slide.shapes.add_textbox(
        Inches(0.38), Inches(7.20), Inches(11.6), Inches(0.22)
    )
    r = tb.text_frame.paragraphs[0].add_run()
    set_run(
        r,
        f"DockSight Improvised Version  ·  Repo3 / V3  ·  {section}",
        10,
        False,
        MUTED,
    )

    nb = slide.shapes.add_textbox(
        Inches(12.25), Inches(7.18), Inches(0.65), Inches(0.24)
    )
    p = nb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    set_run(r, f"{number:02d}", 11, True, CYAN)


def eyebrow(slide, text):
    box = slide.shapes.add_textbox(
        Inches(0.48), Inches(0.20), Inches(11.8), Inches(0.28)
    )
    r = box.text_frame.paragraphs[0].add_run()
    set_run(r, text.upper(), 10, True, CYAN)


def title(slide, text, subtitle=None):
    box = slide.shapes.add_textbox(
        Inches(0.48), Inches(0.52), Inches(12.25), Inches(0.55)
    )
    tf = box.text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    set_run(r, text, 27, True, TEXT)
    if subtitle:
        sub = slide.shapes.add_textbox(
            Inches(0.48), Inches(1.02), Inches(12.15), Inches(0.40)
        )
        sub.text_frame.word_wrap = True
        r = sub.text_frame.paragraphs[0].add_run()
        set_run(r, subtitle, 13, False, MUTED)


def card(slide, left, top, width, height, heading, lines, accent=CYAN, size=13):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = PANEL
    shp.line.color.rgb = LINE

    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(0.07),
        Inches(height),
    )
    fill(bar, accent)

    tb = slide.shapes.add_textbox(
        Inches(left + 0.22),
        Inches(top + 0.14),
        Inches(width - 0.40),
        Inches(height - 0.26),
    )
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    p = tf.paragraphs[0]
    r = p.add_run()
    set_run(r, heading, 14, True, accent)
    for line in lines:
        p = tf.add_paragraph()
        p.space_before = Pt(7)
        r = p.add_run()
        set_run(r, line, size, False, TEXT)


def metric(slide, left, top, width, value, label, color=CYAN):
    box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(1.23),
    )
    box.fill.solid()
    box.fill.fore_color.rgb = PANEL_ALT
    box.line.color.rgb = LINE
    tb = slide.shapes.add_textbox(
        Inches(left + 0.15), Inches(top + 0.12), Inches(width - 0.30), Inches(0.96)
    )
    tf = tb.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    set_run(r, value, 24, True, color)
    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    set_run(r, label, 10, False, MUTED)


def table(slide, left, top, width, rows, col_widths, row_height=0.52, font_size=11):
    shape = slide.shapes.add_table(
        len(rows),
        len(rows[0]),
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(row_height * len(rows)),
    )
    tbl = shape.table
    for i, width_value in enumerate(col_widths):
        tbl.columns[i].width = Inches(width_value)
    for ri, row in enumerate(rows):
        for ci, value in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.text = ""
            cell.margin_left = Inches(0.10)
            cell.margin_right = Inches(0.10)
            cell.margin_top = Inches(0.05)
            cell.margin_bottom = Inches(0.04)
            cell.fill.solid()
            cell.fill.fore_color.rgb = PANEL_ALT if ri == 0 else (
                PANEL if ri % 2 else BG
            )
            p = cell.text_frame.paragraphs[0]
            r = p.add_run()
            set_run(
                r,
                value,
                font_size if ri else font_size + 1,
                ri == 0,
                TEXT if ri else MUTED,
            )
    return shape


def arrow(slide, left, top, width, label, accent=CYAN):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.CHEVRON,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(0.78),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = PANEL_ALT
    shape.line.color.rgb = accent
    tb = slide.shapes.add_textbox(
        Inches(left + 0.08), Inches(top + 0.19), Inches(width - 0.25), Inches(0.30)
    )
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    set_run(r, label, 11, True, TEXT)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 — Cover
    s = prs.slides.add_slide(blank)
    chrome(s, 1, "Stakeholder Presentation")
    eyebrow(s, "Warehouse fulfillment control tower")
    box = s.shapes.add_textbox(
        Inches(0.58), Inches(1.43), Inches(11.8), Inches(1.35)
    )
    tf = box.text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    set_run(r, "DockSight", 44, True, TEXT)
    p = tf.add_paragraph()
    r = p.add_run()
    set_run(r, "Improvised Version", 29, True, CYAN)

    sub = s.shapes.add_textbox(
        Inches(0.62), Inches(3.05), Inches(11.6), Inches(1.10)
    )
    tf = sub.text_frame
    r = tf.paragraphs[0].add_run()
    set_run(r, "Repo3 / V3 — from evidence gate to durable local product simulation", 20, False, TEXT)
    p = tf.add_paragraph()
    p.space_before = Pt(10)
    r = p.add_run()
    set_run(
        r,
        "Three applications · deterministic fulfillment · simulated resources · no physical actuation",
        15,
        False,
        MUTED,
    )

    card(
        s,
        0.62,
        5.08,
        12.02,
        1.50,
        "Presentation boundary",
        [
            "A working local simulation, not a production robot controller.",
            "No AI model, no agent execution, no live WMS/FMS/PLC/robot connection.",
        ],
        AMBER,
        13,
    )
    notes(
        s,
        "Open with the evolution: V2 proved deterministic gates. V3 turns that reasoning into a clickable, persistent local fulfillment product. "
        "Say 'simulation' early. Do not call it a live autonomous warehouse.",
    )

    # 2 — Executive story
    s = prs.slides.add_slide(blank)
    chrome(s, 2, "Executive Story")
    eyebrow(s, "Why V3 exists")
    title(
        s,
        "From “should we proceed?” to “show the complete simulated journey”",
        "V2 established the evidence and refusal logic. V3 productizes intake, accounting, work, recovery and operator workflows.",
    )
    card(
        s,
        0.48,
        1.55,
        3.86,
        4.95,
        "V2 — evidence gate",
        [
            "Observe brownfield contradictions.",
            "Apply deterministic eligibility.",
            "Refuse unsafe or uncertain work.",
            "ActionExecutor remains a stub.",
            "Physical control disabled.",
        ],
        MUTED,
    )
    card(
        s,
        4.73,
        1.55,
        3.86,
        4.95,
        "V3 — product simulation",
        [
            "Bazaar customer order intake.",
            "Transactional reservations and ledger.",
            "Persisted Pick → Move → Pack_feed → Stage.",
            "Supervisor, Fleet and Admin workspaces.",
            "Failure, recovery, charging and replay.",
        ],
        CYAN,
    )
    card(
        s,
        8.97,
        1.55,
        3.88,
        4.95,
        "Still outside scope",
        [
            "Secure production identity and RBAC.",
            "Customer deployment and SLOs.",
            "Real inventory or sensor truth.",
            "External command/write integration.",
            "OM17–21 completion and ROI proof.",
        ],
        RED,
    )
    notes(
        s,
        "The key point is not 'V3 is production.' The key point is 'V3 is materially more complete as a local product simulation.'",
    )

    # 3 — Product at a glance
    s = prs.slides.add_slide(blank)
    chrome(s, 3, "Product")
    eyebrow(s, "What shipped")
    title(
        s,
        "One local origin, three purpose-built applications",
        "A shared FastAPI control plane serves Core Warehouse, FDE Bazaar and the Robot Fleet Simulator.",
    )
    card(
        s,
        0.48,
        1.55,
        3.85,
        3.62,
        "FDE Bazaar",
        [
            "Customer order entry.",
            "Multi-SKU requests and service selection.",
            "Durable HTTP outbox and retry.",
            "Separate delivery vs fulfillment status.",
        ],
        AMBER,
    )
    card(
        s,
        4.74,
        1.55,
        3.85,
        3.62,
        "Core Warehouse",
        [
            "Order queue and progress.",
            "Inventory and resource inspection.",
            "Supervisor, Fleet and Admin workspaces.",
            "Forecast, audit and replay views.",
        ],
        CYAN,
    )
    card(
        s,
        8.98,
        1.55,
        3.86,
        3.62,
        "Robot Fleet Simulator",
        [
            "Current claims and task lifecycle.",
            "Failure, kill, replacement and recovery.",
            "Battery and charging simulation.",
            "No physical robot connection.",
        ],
        GREEN,
    )
    metric(s, 0.48, 5.43, 2.82, "3", "React applications", CYAN)
    metric(s, 3.66, 5.43, 2.82, "1", "FastAPI gateway", CYAN)
    metric(s, 6.85, 5.43, 2.82, "2", "Active SQLite stores", AMBER)
    metric(s, 10.03, 5.43, 2.82, "19", "Inherited source datasets", GREEN)
    notes(
        s,
        "Core is the operational cockpit. Bazaar owns order entry. Fleet Simulator owns resource lifecycle demonstrations. "
        "They share one local gateway but preserve product boundaries.",
    )

    # 4 — End-to-end
    s = prs.slides.add_slide(blank)
    chrome(s, 4, "Journey")
    eyebrow(s, "Customer request to simulated completion")
    title(
        s,
        "A durable, deterministic fulfillment loop",
        "Every transition is persisted; retries and process restarts must not repeat reservations or stock movements.",
    )
    labels = [
        ("1  Create", AMBER),
        ("2  Accept", CYAN),
        ("3  Reserve", CYAN),
        ("4  Assign", GREEN),
        ("5  Execute", GREEN),
        ("6  Observe", AMBER),
    ]
    x = 0.48
    for label, color in labels:
        arrow(s, x, 1.75, 1.98, label, color)
        x += 2.08
    card(
        s,
        0.48,
        2.83,
        4.00,
        3.66,
        "Intake",
        [
            "Bazaar saves order + delivery state first.",
            "Stable request identity protects retries.",
            "Selected warehouse is required for new Bazaar orders.",
        ],
        AMBER,
    )
    card(
        s,
        4.67,
        2.83,
        4.00,
        3.66,
        "Fulfillment",
        [
            "Validate every requested SKU.",
            "Debit free source quantities once.",
            "Credit reservation and create staged work atomically.",
        ],
        CYAN,
    )
    card(
        s,
        8.85,
        2.83,
        4.00,
        3.66,
        "Execution",
        [
            "Acquire eligible, unclaimed simulated resources.",
            "Persist start, due, completion and claims.",
            "Release Pick reservation without a second debit.",
        ],
        GREEN,
    )
    notes(
        s,
        "Stress the consistency boundaries. Intake is durable before handoff. Acceptance and inventory accounting are atomic. "
        "Resource acquisition is deliberately separate so a missing robot does not erase accepted stock.",
    )

    # 5 — Inventory policy
    s = prs.slides.add_slide(blank)
    chrome(s, 5, "Inventory")
    eyebrow(s, "Accounting and policy fork")
    title(
        s,
        "Conservative allocation without claiming physical truth",
        "V3 uses source disagreement operationally, but corrections remain app-owned simulation state.",
    )
    table(
        s,
        0.48,
        1.52,
        12.36,
        [
            ["Event", "Free WMS / ERP / vision", "Reserved", "Picked / in process"],
            ["Order accepted", "− quantity from all three", "+ quantity", "No change"],
            ["Pick starts", "No change", "No change", "No change"],
            ["Pick completes", "No second debit", "− quantity", "+ quantity"],
            ["Cancel before Pick", "+ unpicked quantity", "− unpicked quantity", "No change"],
            ["Cancel after Pick", "No automatic restock", "No picked reservation", "Recovery required"],
        ],
        [3.18, 3.78, 2.55, 2.85],
        row_height=0.57,
        font_size=10,
    )
    card(
        s,
        0.48,
        5.32,
        5.98,
        1.36,
        "Allocation truth",
        ["Available at a location = nonnegative min(WMS, ERP, vision), subject to blockers."],
        GREEN,
        11,
    )
    card(
        s,
        6.68,
        5.32,
        6.16,
        1.36,
        "Important V2 → V3 fork",
        ["Supervisor closure may sync modeled free quantities to their current maximum. This is not physical truth or an external WMS write."],
        AMBER,
        11,
    )
    notes(
        s,
        "This is one of the most important honesty slides. V2 retained the triple as uncertain. V3 adds a modeled correction policy. "
        "Never describe that correction as a site count unless real evidence exists.",
    )

    # 6 — Resource lifecycle
    s = prs.slides.add_slide(blank)
    chrome(s, 6, "Resource Lifecycle")
    eyebrow(s, "Robot and control-asset simulation")
    title(
        s,
        "Available is not the same as eligible, compatible or unclaimed",
        "The scheduler uses explicit rules; failure and recovery preserve completed work.",
    )
    card(
        s,
        0.48,
        1.52,
        3.90,
        4.93,
        "Eligibility",
        [
            "HEALTHY resource state.",
            "Certificate present and not EXPIRED.",
            "ONLINE or INTERMITTENT connectivity.",
            "No OPEN / IN_PROGRESS maintenance.",
            "Warehouse, type and payload compatible.",
        ],
        CYAN,
    )
    card(
        s,
        4.71,
        1.52,
        3.90,
        4.93,
        "Exclusive lifecycle",
        [
            "One active resource claim.",
            "Persisted assignment token.",
            "Configurable simulated duration.",
            "Restart keeps saved due time.",
            "Completed effects are not replayed.",
        ],
        GREEN,
    )
    card(
        s,
        8.94,
        1.52,
        3.90,
        4.93,
        "Failure and energy",
        [
            "Fail or kill the active assignment.",
            "Replace only the unfinished stage.",
            "Explicit recovery clears simulator failure block.",
            "Idle known-battery robot auto-docks below 10%.",
            "Charging stops at 100%.",
        ],
        AMBER,
    )
    notes(
        s,
        "All of this is simulator policy. Synthetic SKU weights, task times and battery behavior are useful for product flows, not evidence of real physical feasibility.",
    )

    # 7 — People and authority
    s = prs.slides.add_slide(blank)
    chrome(s, 7, "People and Authority")
    eyebrow(s, "Persona workspaces")
    title(
        s,
        "The workflow is role-shaped; the identity mechanism is not production security",
        "Repo3 demonstrates who should do what, while `X-Demo-Persona` remains a caller-controlled demo header.",
    )
    table(
        s,
        0.48,
        1.52,
        12.36,
        [
            ["Persona", "Primary responsibility", "Can", "Must not imply"],
            ["Bazaar user", "Create and inspect customer requests", "Submit and retry delivery", "Warehouse task authority"],
            ["Supervisor", "Order flow and inventory exceptions", "Correct scoped evidence; recover work", "Physical ground truth"],
            ["Fleet manager", "Resource readiness and lifecycle", "Repair, fail, kill, recover, charge", "Operational approval"],
            ["Admin", "Policy, audit, config and registration", "Inspect and configure supported demo state", "Secure superuser identity"],
            ["System", "Deterministic acceptance and scheduling", "Apply encoded policies", "AI judgement or physical command"],
        ],
        [1.65, 3.20, 3.76, 3.75],
        row_height=0.70,
        font_size=10,
    )
    card(
        s,
        0.48,
        5.62,
        12.36,
        1.03,
        "Production blocker",
        ["Authentication, deny-by-default authorization and coverage for every mutating route are required before any shared-network use."],
        RED,
        11,
    )
    notes(
        s,
        "Distinguish an authority model from an access-control model. The product has rich role workflows, but the current persona mechanism is not trusted identity.",
    )

    # 8 — Architecture
    s = prs.slides.add_slide(blank)
    chrome(s, 8, "Architecture")
    eyebrow(s, "As-built local modular monolith")
    title(
        s,
        "Three interfaces, one gateway, two mutable stores, immutable evidence",
        "The design optimizes for a portable local demonstration—not multi-instance production scale.",
    )
    card(
        s,
        0.48,
        1.48,
        2.58,
        2.20,
        "Interfaces",
        ["Core Warehouse", "FDE Bazaar", "Fleet Simulator"],
        CYAN,
    )
    arrow(s, 3.35, 2.17, 1.30, "HTTP", CYAN)
    card(
        s,
        4.94,
        1.48,
        3.28,
        2.20,
        "FastAPI gateway",
        ["Fulfillment", "Personas", "Bazaar", "Lifecycle", "Replay / Review"],
        GREEN,
        11,
    )
    arrow(s, 8.48, 2.17, 1.30, "SQL", GREEN)
    card(
        s,
        10.05,
        1.48,
        2.79,
        2.20,
        "Mutable state",
        ["fulfillment.sqlite", "bazaar.sqlite"],
        AMBER,
    )
    card(
        s,
        0.48,
        4.02,
        4.00,
        2.34,
        "Immutable evidence plane",
        ["CSV / JSONL / legacy SQLite", "Original contradictions remain reviewable."],
        MUTED,
    )
    card(
        s,
        4.67,
        4.02,
        4.00,
        2.34,
        "Runtime controls",
        ["Transactions · revisions · tokens · claims", "Background executor, outbox and forecast workers."],
        GREEN,
    )
    card(
        s,
        8.85,
        4.02,
        4.00,
        2.34,
        "Deployment boundary",
        ["Python launcher binds to 127.0.0.1.", "No TLS, HA, tenant isolation or production IAM."],
        RED,
    )
    notes(
        s,
        "The architecture is deliberately simple and portable. Do not position single-process SQLite as a production scaling decision.",
    )

    # 9 — Forecast and replay
    s = prs.slides.add_slide(blank)
    chrome(s, 9, "Decision Support")
    eyebrow(s, "Forecasting and replay")
    title(
        s,
        "Evidence-backed decision support—without pretending it is AI",
        "Repo3 adds transparent demand baselines and copied-state replay for explanation and comparison.",
    )
    card(
        s,
        0.48,
        1.54,
        5.98,
        4.72,
        "Demand forecast",
        [
            "Mean ordered units over completed IST calendar days.",
            "Includes zero-demand days.",
            "Ceiling-rounded 7-day and 30-day totals.",
            "Explicit no-history and short-history coverage.",
            "Saved evidence stays attached to its run.",
            "Statistical baseline—not a trained model.",
        ],
        CYAN,
    )
    card(
        s,
        6.70,
        1.54,
        6.14,
        4.72,
        "Historical replay",
        [
            "Copy live snapshot into an isolated simulation DB.",
            "Run policy without mutating the operational demo store.",
            "Compare simulation outputs with archive benchmarks.",
            "Retain manifests and CSV reports.",
            "Descriptive comparison—not causal value proof.",
            "Some replay payloads are absent from the supplied ZIP.",
        ],
        AMBER,
    )
    notes(
        s,
        "This slide protects the AI claim. The forecast is useful and transparent, but it is not ML. Replay supports explanation; it does not prove production resilience or ROI.",
    )

    # 10 — Evidence
    s = prs.slides.add_slide(blank)
    chrome(s, 10, "Evidence")
    eyebrow(s, "What is proven and at what confidence")
    title(
        s,
        "Separate as-built evidence from packaged test claims",
        "The V3 pack records direct inspection, fresh checks and prior package verification as different evidence levels.",
    )
    metric(s, 0.48, 1.53, 2.76, "73", "isolated regressions reported", GREEN)
    metric(s, 3.63, 1.53, 2.76, "25", "focused package tests reported", GREEN)
    metric(s, 6.79, 1.53, 2.76, "516", "present files matching SHA-256", CYAN)
    metric(s, 9.94, 1.53, 2.90, "45", "manifest files absent", AMBER)
    card(
        s,
        0.48,
        3.12,
        3.88,
        3.14,
        "Direct inspection",
        [
            "Runtime, APIs, databases and boundaries reviewed.",
            "Python compile check passed.",
        ],
        CYAN,
    )
    card(
        s,
        4.72,
        3.12,
        3.88,
        3.14,
        "Packaged report",
        [
            "Linux builds, browser checks and DB integrity reported.",
            "Treat as prior evidence, not a fresh Windows rerun.",
        ],
        GREEN,
    )
    card(
        s,
        8.96,
        3.12,
        3.88,
        3.14,
        "Still required",
        [
            "Install pytest and pnpm on a clean Windows target.",
            "Rerun backend, typecheck/build, Playwright and DB checks.",
        ],
        AMBER,
    )
    notes(
        s,
        "Do not blend all test counts into one claim. Be explicit: these are historical reports plus a fresh compile/hash inspection. "
        "The missing files are backup/replay companions, with zero mismatches among present files.",
    )

    # 11 — FDE status
    s = prs.slides.add_slide(blank)
    chrome(s, 11, "FDE Operating Model")
    eyebrow(s, "OM1–21 status")
    title(
        s,
        "Repo3 strengthens product engineering—not customer lifecycle proof",
        "V2 remains the source of truth for discovery, invariants, hard gates and original evals.",
    )
    card(
        s,
        0.48,
        1.53,
        3.88,
        4.98,
        "OM1–8",
        [
            "Inherited from the separate V2 evidence spine.",
            "Repo3 does not redo field immersion, RCA or option selection.",
            "Embedded brownfield folder is a trimmed baseline—not full V2.",
        ],
        MUTED,
    )
    card(
        s,
        4.72,
        1.53,
        3.88,
        4.98,
        "OM9–16",
        [
            "Strong local product engineering.",
            "Information architecture, application, tests, packaging and recovery.",
            "Security and independent assurance remain demo-limited.",
        ],
        GREEN,
    )
    card(
        s,
        8.96,
        1.53,
        3.88,
        4.98,
        "OM17–21",
        [
            "Local ZIP is not progressive customer deployment.",
            "Replay is not production monitoring.",
            "7,000+ simulated orders are not ROI.",
            "No AIMS or governed retirement evidence.",
        ],
        RED,
    )
    notes(
        s,
        "Say plainly: OM17–21 are still partial or not proven. Product completeness in a local simulator is not the same as customer lifecycle completion.",
    )

    # 12 — Release posture
    s = prs.slides.add_slide(blank)
    chrome(s, 12, "Release")
    eyebrow(s, "Readiness and risk")
    title(
        s,
        "Conditional local-demo GO; shared or physical deployment NO-GO",
        "The current trust boundary is loopback plus synthetic data.",
    )
    card(
        s,
        0.48,
        1.52,
        3.88,
        4.98,
        "Local demonstration",
        [
            "MAY release after target-machine test rerun.",
            "Bind to loopback.",
            "Use synthetic data only.",
            "Verify active databases and package scope.",
        ],
        GREEN,
    )
    card(
        s,
        4.72,
        1.52,
        3.88,
        4.98,
        "Shared network",
        [
            "NO-GO today.",
            "Needs IAM, RBAC, CSRF/TLS, secrets and route-level authorization.",
            "Needs migrations, observability, SLOs, load and security testing.",
        ],
        AMBER,
    )
    card(
        s,
        8.96,
        1.52,
        3.88,
        4.98,
        "Physical / external write",
        [
            "PROHIBITED in this mandate.",
            "Needs verified interfaces, command IDs and outcome protocol.",
            "Needs hazard analysis, safety case, shadow mode and independent TEVV.",
        ],
        RED,
    )
    notes(
        s,
        "This is the release decision. Avoid soft language: local demo can proceed conditionally; shared network and physical integration cannot.",
    )

    # 13 — Value
    s = prs.slides.add_slide(blank)
    chrome(s, 13, "Value")
    eyebrow(s, "What to measure next")
    title(
        s,
        "V3 creates measurable operational data—but not proven business value yet",
        "Targets and a live comparison plan are still required before benefits or ROI claims.",
    )
    table(
        s,
        0.48,
        1.52,
        12.36,
        [
            ["Measure", "Available from V3", "What is still missing"],
            ["Intake integrity", "Duplicate/retry identities and delivery state", "Target error rate and live volume"],
            ["Stock accounting", "Reservations and movement ledger", "Independent reconciliation to real inventory"],
            ["Flow", "Order, stage, hold and completion timestamps", "Approved lead-time and aging targets"],
            ["Recovery", "Failure, replacement and intervention history", "Real incident baseline and MTTR definition"],
            ["Forecast", "Saved 7/30-day statistical baseline", "Backtest error target and source completeness"],
            ["Value", "Descriptive simulated outcomes", "Controlled before/after and financial inputs"],
        ],
        [2.65, 4.55, 5.16],
        row_height=0.66,
        font_size=10,
    )
    notes(
        s,
        "Avoid turning the 7,009-order snapshot or the archive on-time percentage into benefits proof. "
        "The product now emits useful measurement data; the business targets and live treatment comparison do not yet exist.",
    )

    # 14 — Demo and ask
    s = prs.slides.add_slide(blank)
    chrome(s, 14, "Close")
    eyebrow(s, "Recommended demonstration and decision")
    title(
        s,
        "Show one order, one hold, one recovery—and keep the boundary visible",
        "The ask is approval of the local product simulation and its next verification step, not approval for live control.",
    )
    card(
        s,
        0.48,
        1.52,
        3.88,
        4.92,
        "Demo sequence",
        [
            "1. Create a Bazaar order.",
            "2. Show atomic acceptance or shortage rejection.",
            "3. Watch staged simulated execution.",
            "4. Fail/replace one assigned resource.",
            "5. Inspect audit, forecast and progress.",
        ],
        CYAN,
    )
    card(
        s,
        4.72,
        1.52,
        3.88,
        4.92,
        "Decision requested",
        [
            "Accept V3 as the updated local Repo3 product baseline.",
            "Approve clean-host Windows verification.",
            "Retain deterministic/no-AI/no-OT architecture.",
            "Keep V2 as the original FDE evidence spine.",
        ],
        GREEN,
    )
    card(
        s,
        8.96,
        1.52,
        3.88,
        4.92,
        "Do not approve yet",
        [
            "Shared network exposure.",
            "Production multi-user rollout.",
            "External WMS/FMS/PLC integration.",
            "Physical robot command path.",
            "Modernization-complete or ROI claim.",
        ],
        RED,
    )
    notes(
        s,
        "Close with the decision: DockSight V3 is a credible, durable local simulator and stakeholder demo. "
        "The next responsible step is clean-host verification and explicit pilot architecture—not turning on a live write path.",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
