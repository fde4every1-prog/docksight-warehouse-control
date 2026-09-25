"""Bounded XLSX ingestion and normalized transfer review."""
from __future__ import annotations

import io
import math
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from typing import Any

from openpyxl import Workbook, load_workbook

HEADERS = ["Order ID", "SKU", "Qty", "Order Cut-off"]
MAX_BYTES = 2 * 1024 * 1024
MAX_ROWS = 500


class WorkbookError(ValueError):
    def __init__(self, code: str, message: str, errors: list[dict[str, Any]] | None = None):
        super().__init__(message)
        self.code, self.message, self.errors = code, message, errors or []


def _parse_cutoff(value: Any) -> tuple[str | None, str | None]:
    if value in (None, ""):
        return None, None
    if isinstance(value, datetime):
        parsed = value
    else:
        try:
            parsed = datetime.fromisoformat(str(value).strip().replace("Z", "+00:00"))
        except (TypeError, ValueError):
            return None, "Order cut-off must be an RFC 3339 timestamp with timezone."
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        return None, "Order cut-off must include an explicit timezone."
    return parsed.astimezone(timezone.utc).isoformat().replace("+00:00", "Z"), None


def parse_xlsx(data: bytes) -> list[dict[str, Any]]:
    if len(data) > MAX_BYTES:
        raise WorkbookError("FILE_TOO_LARGE", "Workbook exceeds the 2 MiB limit.")
    if not data.startswith(b"PK"):
        raise WorkbookError("INVALID_XLSX", "Body is not an XLSX workbook.")
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise WorkbookError("INVALID_XLSX", "Workbook could not be read as XLSX.") from exc
    try:
        if len(book.worksheets) != 1:
            raise WorkbookError("SHEET_LIMIT", "Workbook must contain exactly one worksheet.")
        sheet = book.worksheets[0]
        if sheet.max_row - 1 > MAX_ROWS:
            raise WorkbookError("ROW_LIMIT", f"Workbook has more than {MAX_ROWS} data rows.")
        header = [sheet.cell(1, index + 1).value for index in range(4)]
        if header != HEADERS or sheet.max_column != 4:
            raise WorkbookError("INVALID_HEADERS", f"Headers must be exactly: {', '.join(HEADERS)}.")
        rows = []
        for number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if all(v in (None, "") for v in values):
                continue
            if len(rows) >= MAX_ROWS:
                raise WorkbookError("ROW_LIMIT", f"Workbook has more than {MAX_ROWS} populated rows.")
            rows.append({
                "row_id": f"ROW-{number}", "source_row": number,
                "order_id": str(values[0]).strip() if values[0] not in (None, "") else None,
                "sku": str(values[1]).strip() if values[1] not in (None, "") else "",
                "quantity_raw": values[2], "order_cutoff_raw": values[3],
            })
        if not rows:
            raise WorkbookError("EMPTY_WORKBOOK", "Workbook must contain at least one data row.")
        return rows
    finally:
        book.close()


def sample_rows(snapshot: dict[str, Any]) -> list[dict[str, Any]]:
    candidates = [
        row for row in snapshot["inventory"]
        if row["available_unreserved_qty"] > 0 and isinstance(row.get("weight_kg"), (int, float))
        and float(row["weight_kg"]) > 0
    ]
    # Keep the starter scenario comfortably feasible for the selected live
    # actors while retaining authentic SKU/bin/weight evidence.
    max_transfer = max(
        (float(robot["capacity_kg"]) for robot in snapshot["robots"]
         if robot["id"] in snapshot["selected_robot_ids"] and robot["eligible"]
         and "transfer" in robot["capabilities"] and isinstance(robot.get("capacity_kg"), (int, float))),
        default=0,
    )
    candidates = sorted(
        [row for row in candidates if float(row["weight_kg"]) <= max_transfer],
        key=lambda row: (float(row["weight_kg"]), str(row["sku"]), str(row["location"])),
    )
    unique = []
    seen = set()
    for row in candidates:
        if row["sku"] not in seen:
            unique.append(row)
            seen.add(row["sku"])
        if len(unique) == 3:
            break
    candidates = unique
    if not candidates:
        raise WorkbookError("NO_ELIGIBLE_STOCK", "Snapshot has no eligible weighted stock for an example.")
    clock = datetime.fromisoformat(snapshot["clock"]["at"].replace("Z", "+00:00"))
    result = []
    for index, row in enumerate(candidates, 1):
        result.append({
            "order_id": f"EXAMPLE-{index:02d}", "sku": row["sku"], "qty": 1,
            "order_cutoff": (clock + timedelta(minutes=20 + index * 5)).isoformat().replace("+00:00", "Z"),
            "example_only": True, "example_bin": row["location"],
        })
    return result


def template_bytes(snapshot: dict[str, Any]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Transfer Orders"
    sheet.append(HEADERS)
    for row in sample_rows(snapshot):
        sheet.append([row["order_id"], row["sku"], row["qty"], row["order_cutoff"]])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:D{sheet.max_row}"
    for column, width in {"A": 22, "B": 22, "C": 10, "D": 30}.items():
        sheet.column_dimensions[column].width = width
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def normalize(raw_rows: list[dict[str, Any]], snapshot: dict[str, Any]) -> tuple[list[dict[str, Any]], list[str]]:
    by_sku: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for stock in snapshot["inventory"]:
        by_sku[str(stock["sku"])].append(stock)
    remaining = {str(stock["source_id"]): int(stock["available_unreserved_qty"]) for stock in snapshot["inventory"]}
    now = datetime.fromisoformat(snapshot["clock"]["at"].replace("Z", "+00:00"))
    rows = []
    pair_counts = Counter(
        (str(r.get("order_id") or ""), str(r.get("sku") or ""))
        for r in raw_rows if r.get("order_id") and r.get("sku")
    )
    cutoffs: dict[str, set[str]] = defaultdict(set)
    parsed_cutoffs: dict[str, str | None] = {}
    for raw in raw_rows:
        cutoff, _ = _parse_cutoff(raw.get("order_cutoff_raw"))
        parsed_cutoffs[raw["row_id"]] = cutoff
        if raw.get("order_id") and cutoff:
            cutoffs[str(raw["order_id"])].add(cutoff)
    for raw in raw_rows:
        errors: list[dict[str, str]] = []
        def error(code: str, field: str, message: str) -> None:
            errors.append({"code": code, "field": field, "message": message})
        oid, sku = raw.get("order_id"), str(raw.get("sku") or "").strip()
        if not oid:
            error("MISSING_ORDER_ID", "order_id", "Order ID is required before planning.")
        if not sku:
            error("MISSING_SKU", "sku", "SKU is required.")
        quantity = None
        value = raw.get("quantity_raw")
        if isinstance(value, bool):
            pass
        elif isinstance(value, int):
            quantity = value
        elif isinstance(value, float) and value.is_integer():
            quantity = int(value)
        elif isinstance(value, str) and value.strip().isdigit():
            quantity = int(value.strip())
        if quantity is None or quantity <= 0:
            error("INVALID_QUANTITY", "quantity", "Qty must be a positive integer.")
        cutoff, cutoff_error = _parse_cutoff(raw.get("order_cutoff_raw"))
        if cutoff_error:
            error("INVALID_CUTOFF", "order_cutoff", cutoff_error)
        elif not cutoff:
            error("MISSING_CUTOFF", "order_cutoff", "Order cut-off is required before planning.")
        elif datetime.fromisoformat(cutoff.replace("Z", "+00:00")) <= now:
            error("CUTOFF_NOT_FUTURE", "order_cutoff", "Order cut-off must be after the frozen simulation clock.")
        if oid and len(cutoffs[str(oid)]) > 1:
            error("CONFLICTING_CUTOFF", "order_cutoff", "All lines in an order must have the same cut-off.")
        if oid and sku and pair_counts[(str(oid), sku)] > 1:
            error("DUPLICATE_LINE", "sku", "Duplicate Order ID + SKU rows must be combined.")
        stocks = by_sku.get(sku, [])
        if sku and not stocks:
            error("UNKNOWN_SKU", "sku", f"{sku} has no DC-01 inventory evidence.")
        valid_weights = {
            float(stock["weight_kg"]) for stock in stocks
            if isinstance(stock.get("weight_kg"), (int, float))
            and math.isfinite(float(stock["weight_kg"])) and float(stock["weight_kg"]) > 0
        }
        if stocks and (len(valid_weights) != 1 or any(stock.get("weight_kg") is None for stock in stocks)):
            error("MISSING_OR_CONFLICTING_WEIGHT", "sku", f"{sku} lacks one consistent positive unit weight.")
        allocations = []
        needed = quantity or 0
        if quantity and stocks:
            for stock in sorted(stocks, key=lambda item: (item["location"], item["source_id"])):
                available = remaining.get(str(stock["source_id"]), 0)
                take = min(needed, available)
                if take:
                    allocations.append({
                        "inventory_source_id": stock["source_id"], "bin_id": stock["location"],
                        "zone_id": stock["zone"], "quantity": take,
                        "node_id": stock["zone"], "mapping": "live_zone_id_assumed_zone_center",
                    })
                    remaining[str(stock["source_id"])] -= take
                    needed -= take
                if needed == 0:
                    break
            if needed:
                available = sum(int(stock["available_unreserved_qty"]) for stock in stocks)
                error("STOCK_SHORTAGE", "quantity", f"{sku} requests {quantity}; frozen available stock is {available}.")
        weight = next(iter(valid_weights), None) if len(valid_weights) == 1 else None
        rows.append({
            "row_id": raw["row_id"], "source_row": raw["source_row"], "order_id": oid,
            "sku": sku, "quantity": quantity, "order_cutoff": cutoff, "cutoff_timezone": "UTC" if cutoff else None,
            "allocations": allocations, "unit_weight_kg": weight,
            "total_weight_kg": weight * quantity if weight and quantity else None,
            "errors": errors, "warnings": [
                {"code": "ASSUMED_COORDINATE", "field": "allocations",
                 "message": "Live bin/zone provenance retained; displayed zone-center coordinates are illustrative."}
            ] if allocations else [],
            "provenance": {"source": "xlsx", "snapshot_id": snapshot["snapshot_id"]},
        })
    blockers = [f"{row['row_id']}: {error['message']}" for row in rows for error in row["errors"]]
    return rows, blockers