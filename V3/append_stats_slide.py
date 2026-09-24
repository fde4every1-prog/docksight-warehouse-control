"""Append the single 'engagement in numbers' stats slide.

Run AFTER the deck and the KPI slides exist:

    python generate_repo3_solution_deck.py        # slides 1-13
    python append_kpi_comparison_slides.py        # slides 14-16
    python append_stats_slide.py                  # slide 17

Numbers are counted from the V2 and V3 trees:
    V2 spine   53 = 22 discovery + 21 specs + 2 spec diagrams + 5 evals + 3 glue
    V3 pack    17 = 1 PRD + 6 ADRs + 9 assurance docs + 1 README
    ADRs        8 = 2 V2 + 6 V3
    invariants 26 = 12 V2 (I1-I12) + 14 V3 (V3-I1-V3-I14)
    scripts    10 = 6 V2 scripts + 4 V3 scripts
"""

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from generate_repo3_solution_deck import (
    AMBER,
    CHARCOAL,
    DARK_GRAY,
    GREEN,
    LIGHT_GRAY,
    MEDIUM_GRAY,
    WHITE,
    YELLOW,
    base,
    fill,
    heading,
    notes,
    set_run,
)

HERE = Path(__file__).resolve().parent
DECK = HERE / "DockSight_Repo3_Solution_Deck.pptx"


def hero(slide, left, top, width, value, unit, label, accent):
    """Large headline number tile."""
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(1.04),
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = DARK_GRAY
    shp.line.color.rgb = accent

    tb = slide.shapes.add_textbox(
        Inches(left + 0.10), Inches(top + 0.08),
        Inches(width - 0.20), Inches(0.52),
    )
    tf = tb.text_frame
    tf.margin_left = 0
    tf.margin_top = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    set_run(r, value, 30, True, accent)
    if unit:
        r = p.add_run()
        set_run(r, " " + unit, 11, False, LIGHT_GRAY)

    lb = slide.shapes.add_textbox(
        Inches(left + 0.10), Inches(top + 0.62),
        Inches(width - 0.20), Inches(0.34),
    )
    ltf = lb.text_frame
    ltf.word_wrap = True
    ltf.margin_left = 0
    ltf.margin_top = 0
    p = ltf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    set_run(r, label, 9, False, LIGHT_GRAY)


def mini(slide, left, top, width, count, label, accent):
    """Compact key-artifact panel."""
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(1.00),
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = DARK_GRAY
    shp.line.color.rgb = MEDIUM_GRAY

    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left), Inches(top), Inches(0.06), Inches(1.00),
    )
    fill(bar, accent)

    tb = slide.shapes.add_textbox(
        Inches(left + 0.16), Inches(top + 0.09),
        Inches(width - 0.30), Inches(0.34),
    )
    tf = tb.text_frame
    tf.margin_left = 0
    tf.margin_top = 0
    r = tf.paragraphs[0].add_run()
    set_run(r, count, 17, True, accent)

    lb = slide.shapes.add_textbox(
        Inches(left + 0.16), Inches(top + 0.45),
        Inches(width - 0.30), Inches(0.48),
    )
    ltf = lb.text_frame
    ltf.word_wrap = True
    ltf.margin_left = 0
    ltf.margin_top = 0
    r = ltf.paragraphs[0].add_run()
    set_run(r, label, 9, False, LIGHT_GRAY)


def stage(slide, left, top, width, height, title, count, detail, accent):
    """Big-picture stage block."""
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height),
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = DARK_GRAY
    shp.line.color.rgb = accent

    tb = slide.shapes.add_textbox(
        Inches(left + 0.18), Inches(top + 0.10),
        Inches(width - 0.36), Inches(0.28),
    )
    tb.text_frame.margin_left = 0
    tb.text_frame.margin_top = 0
    r = tb.text_frame.paragraphs[0].add_run()
    set_run(r, title, 12, True, accent)

    cb = slide.shapes.add_textbox(
        Inches(left + 0.18), Inches(top + 0.40),
        Inches(width - 0.36), Inches(0.34),
    )
    cb.text_frame.margin_left = 0
    cb.text_frame.margin_top = 0
    r = cb.text_frame.paragraphs[0].add_run()
    set_run(r, count, 19, True, WHITE)

    db = slide.shapes.add_textbox(
        Inches(left + 0.18), Inches(top + 0.76),
        Inches(width - 0.36), Inches(0.40),
    )
    dtf = db.text_frame
    dtf.word_wrap = True
    dtf.margin_left = 0
    dtf.margin_top = 0
    r = dtf.paragraphs[0].add_run()
    set_run(r, detail, 9, False, LIGHT_GRAY)


def chevron(slide, left, top, color):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW, Inches(left), Inches(top), Inches(0.42), Inches(0.24)
    )
    fill(shp, color)


def coverage_bar(slide, left, top, width, height, done, total, accent_done, accent_rest):
    done_w = width * done / total
    a = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(done_w), Inches(height)
    )
    fill(a, accent_done)
    b = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left + done_w), Inches(top), Inches(width - done_w), Inches(height),
    )
    fill(b, accent_rest)

    for shape, text in ((a, "OM 1-16 evidenced"), (b, "OM 17-21 deferred")):
        tf = shape.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.08)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        set_run(r, text, 10, True, CHARCOAL)


def build(prs, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    base(slide, number, "Engagement in Numbers", dark=True)
    heading(
        slide,
        "Engagement in numbers",
        "One 21-step FDE spine, 70 artifacts, one working product",
        dark=True,
    )

    # Row A - headline numbers
    xs = (0.48, 2.97, 5.46, 7.95, 10.44)
    hero(slide, xs[0], 1.46, 2.41, "70", "", "FDE artifacts created", YELLOW)
    hero(slide, xs[1], 1.46, 2.41, "21", "", "operating-model workflows followed", YELLOW)
    hero(slide, xs[2], 1.46, 2.41, "8", "", "architecture decisions on record", GREEN)
    hero(slide, xs[3], 1.46, 2.41, "26", "", "invariants defined and tested", GREEN)
    hero(slide, xs[4], 1.46, 2.41, "10", "", "scripts to rebuild the evidence", AMBER)

    # Row B - key artifacts
    panel_w = 1.97
    step = 2.08
    panels = [
        ("2", "Charters\nengagement + 90-day", YELLOW),
        ("1", "PRD\n10 FR / 10 NFR", YELLOW),
        ("8", "ADRs\nincl. LLM off write path", GREEN),
        ("2", "Language packs\n12 + 14 invariants", GREEN),
        ("15", "Test suites\ngates and evals", AMBER),
        ("3", "Release gates\nGO / NO-GO / prohibited", AMBER),
    ]
    for index, (count, label, accent) in enumerate(panels):
        mini(slide, 0.48 + index * step, 2.70, panel_w, count, label, accent)

    # Row C - big picture
    stage(
        slide, 0.48, 3.92, 3.84, 1.26,
        "V2 - FDE proof",
        "53 artifacts",
        "Brownfield discovery, domain model, gates, evals. Observe and refuse.",
        YELLOW,
    )
    chevron(slide, 4.44, 4.44, WHITE)
    stage(
        slide, 4.98, 3.92, 3.84, 1.26,
        "Repo 3 - working product",
        "Local simulator",
        "Order intake, reservations, task execution, personas, replay.",
        GREEN,
    )
    chevron(slide, 8.94, 4.44, WHITE)
    stage(
        slide, 9.48, 3.92, 3.37, 1.26,
        "V3 - evidence pack",
        "17 artifacts",
        "PRD, architecture, traceability, readiness, ADRs.",
        AMBER,
    )

    # Row D - OM coverage
    lb = slide.shapes.add_textbox(Inches(0.48), Inches(5.34), Inches(12.37), Inches(0.24))
    lb.text_frame.margin_left = 0
    r = lb.text_frame.paragraphs[0].add_run()
    set_run(r, "OPERATING-MODEL COVERAGE", 9, True, MEDIUM_GRAY)
    coverage_bar(slide, 0.48, 5.60, 12.37, 0.38, 16, 21, YELLOW, MEDIUM_GRAY)

    # Row E - closing statement
    band = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.48), Inches(6.20), Inches(12.37), Inches(0.52),
    )
    band.fill.solid()
    band.fill.fore_color.rgb = DARK_GRAY
    band.line.color.rgb = GREEN
    tb = slide.shapes.add_textbox(
        Inches(0.64), Inches(6.29), Inches(12.05), Inches(0.36)
    )
    tf = tb.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = True
    tf.margin_left = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    set_run(r, "Demo-ready on loopback   ·   ", 11, False, LIGHT_GRAY)
    r = p.add_run()
    set_run(r, "Production NO-GO", 11, True, AMBER)
    r = p.add_run()
    set_run(r, "   ·   ", 11, False, LIGHT_GRAY)
    r = p.add_run()
    set_run(r, "Physical control PROHIBITED", 11, True, AMBER)

    notes(
        slide,
        "This is the closing stats slide. Lead with the two numbers that matter: we followed a "
        "21-step operating model and produced 70 artifacts against it.\n\n"
        "The top row is the headline. Seventy artifacts, twenty-one workflows, eight architecture "
        "decisions on record, twenty-six invariants, and ten scripts that regenerate the decks and "
        "re-verify the numbers. That last one matters more than it looks: the story can be rebuilt "
        "from the repository rather than recalled from memory.\n\n"
        "The second row is the key artifact set. Two charters, so scope and the enhancement window "
        "were written before build. One PRD with ten functional and ten non-functional requirements, "
        "reconstructed from the system as built. Eight ADRs, including the decision to keep the LLM "
        "off the write path. Two ubiquitous-language packs carrying twelve and fourteen invariants. "
        "Fifteen test suites. Three release gates.\n\n"
        "The third row is the bird's-eye view: V2 proved the brownfield can be reconciled and gated "
        "with 53 artifacts, Repo 3 turned that into a working local simulator, and the V3 pack added "
        "17 artifacts so it can be reviewed as a product.\n\n"
        "The coverage bar is the honesty check. Sixteen of the twenty-one workflows are evidenced. "
        "Five - deploy, monitor, value, AIMS, retire - are deferred in writing, because there is no "
        "live customer deployment. We did not fake them.\n\n"
        "Close on the bottom band. Demo-ready on loopback, production NO-GO, physical control "
        "prohibited. Say it plainly. Those boundaries are why the rest of the numbers can be "
        "trusted.",
    )
    return slide


def main():
    prs = Presentation(str(DECK))
    start = len(prs.slides)
    if start != 16:
        raise SystemExit(
            "Expected 16 slides before appending the stats slide, found %d." % start
        )
    build(prs, start + 1)
    prs.save(str(DECK))
    print("appended slide %d; total %d" % (start + 1, len(prs.slides)))


if __name__ == "__main__":
    main()
