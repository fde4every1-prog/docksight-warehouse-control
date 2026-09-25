"""Working deck: always write DockSight_Sample_Slide.pptx. Do not create extra pptx files."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = (
    Path(__file__).resolve().parents[1] / "docs" / "DockSight_Sample_Slide.pptx"
)

PAPER = RGBColor(0xF3, 0xF0, 0xE8)
CARD = RGBColor(0xFF, 0xFC, 0xF7)
INK = RGBColor(0x2A, 0x32, 0x38)
MUTED = RGBColor(0x5A, 0x63, 0x6A)
STEEL = RGBColor(0x3C, 0x58, 0x6A)
STEEL_DEEP = RGBColor(0x2E, 0x44, 0x52)
AMBER = RGBColor(0xD0, 0x9A, 0x24)
AMBER_SOFT = RGBColor(0xF3, 0xE6, 0xC0)
LINE = RGBColor(0xD4, 0xCC, 0xBA)
RACK = RGBColor(0xC8, 0xC0, 0xB0)
STOP = RGBColor(0xA8, 0x4B, 0x40)
GO = RGBColor(0x3E, 0x6F, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _run(run, text, size=16, bold=False, color=INK):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _back(slide, shape):
    tree = slide.shapes._spTree
    el = shape._element
    tree.remove(el)
    tree.insert(2, el)


def _rect(slide, l, t, w, h, color, line=None):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
    return shp


def _round(slide, l, t, w, h, color, line=None):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
    try:
        shp.adjustments[0] = 0.12
    except Exception:
        pass
    return shp


def _oval(slide, l, t, w, h, color):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    _fill(shp, color)
    return shp


def _warehouse_backdrop(slide):
    bg = _rect(slide, 0, 0, 13.333, 7.5, PAPER)
    _back(slide, bg)
    for x in (11.55, 12.15, 12.75):
        _rect(slide, x, 0.55, 0.045, 6.35, RACK)
    for y in (1.35, 2.55, 3.75, 4.95, 6.15):
        _rect(slide, 11.52, y, 1.32, 0.035, RACK)
    _rect(slide, 0, 7.12, 13.333, 0.10, AMBER)
    _rect(slide, 0, 7.22, 13.333, 0.28, STEEL_DEEP)
    _rect(slide, 0, 0, 0.10, 7.5, STEEL)


def _bay_sign(slide):
    _round(slide, 0.42, 0.18, 2.15, 0.38, STEEL)
    tb = slide.shapes.add_textbox(Inches(0.48), Inches(0.22), Inches(2.05), Inches(0.30))
    r = tb.text_frame.paragraphs[0].add_run()
    _run(r, "BAY A   DC-01", 12, True, WHITE)


def _amr(slide, l, t):
    _round(slide, l + 0.18, t, 0.72, 0.12, STEEL)
    _round(slide, l, t + 0.14, 1.08, 0.62, STEEL)
    _rect(slide, l + 0.12, t + 0.28, 0.28, 0.16, AMBER)
    _rect(slide, l + 0.48, t + 0.32, 0.46, 0.08, STEEL_DEEP)
    _oval(slide, l + 0.06, t + 0.68, 0.28, 0.28, STEEL_DEEP)
    _oval(slide, l + 0.74, t + 0.68, 0.28, 0.28, STEEL_DEEP)
    _oval(slide, l + 0.12, t + 0.74, 0.16, 0.16, AMBER_SOFT)
    _oval(slide, l + 0.80, t + 0.74, 0.16, 0.16, AMBER_SOFT)


def _footer(slide, text):
    tb = slide.shapes.add_textbox(Inches(0.42), Inches(7.24), Inches(12.4), Inches(0.22))
    r = tb.text_frame.paragraphs[0].add_run()
    _run(r, text, 11, False, PAPER)


def _chrome(slide, footer):
    _warehouse_backdrop(slide)
    _bay_sign(slide)
    _amr(slide, 11.85, 0.22)
    _footer(slide, footer)


def _card(slide, l, t, w, h, heading, lines, accent, body=13):
    _round(slide, l, t, w, h, CARD, LINE)
    _rect(slide, l, t, 0.08, h, accent)
    tb = slide.shapes.add_textbox(
        Inches(l + 0.22), Inches(t + 0.14), Inches(w - 0.36), Inches(h - 0.26)
    )
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    _run(r, heading, 14, True, accent)
    for line in lines:
        p = tf.add_paragraph()
        p.space_before = Pt(8)
        r = p.add_run()
        _run(r, line, body, False, INK)


def _add_theme_sample(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _chrome(s, "DockSight  ·  working deck  ·  theme sample")

    title = s.shapes.add_textbox(Inches(2.72), Inches(0.16), Inches(8.8), Inches(0.42))
    r = title.text_frame.paragraphs[0].add_run()
    _run(r, "DockSight  ·  sample theme", 22, True, INK)

    sub = s.shapes.add_textbox(Inches(0.42), Inches(0.68), Inches(10.8), Inches(0.36))
    r = sub.text_frame.paragraphs[0].add_run()
    _run(
        r,
        "Light warehouse floor  ·  steel racks  ·  aisle tape  ·  one AMR.",
        13,
        False,
        MUTED,
    )

    _card(
        s,
        0.42,
        1.22,
        4.05,
        5.55,
        "Look",
        [
            "Paper floor, not dark navy.",
            "Yellow tape like a warehouse aisle.",
            "Steel bay sign and rack lines.",
            "One small robot — not a collage.",
            "Cards sit on the floor like pallets.",
        ],
        STEEL,
    )
    _card(
        s,
        4.62,
        1.22,
        4.05,
        5.55,
        "How we will use it",
        [
            "You send the prompt for each slide.",
            "We keep this chrome and swap the words.",
            "No extra icons unless you ask.",
            "Notes go under the slide for you to speak.",
        ],
        AMBER,
    )
    _card(
        s,
        8.82,
        1.22,
        4.08,
        5.55,
        "Placeholder copy",
        [
            "Problem: systems disagree. People still hit Go.",
            "Solution: watch-and-stop before assign.",
            "This box is only to judge type size.",
        ],
        GO,
    )
    s.notes_slide.notes_text_frame.text = "Theme sample. Next slide is the problem statement."


def _add_problem_horizontal(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _chrome(s, "DockSight  ·  working deck  ·  problem statement")

    title = s.shapes.add_textbox(Inches(2.72), Inches(0.16), Inches(8.8), Inches(0.42))
    r = title.text_frame.paragraphs[0].add_run()
    _run(r, "The problem we are here to stop", 22, True, INK)

    sub = s.shapes.add_textbox(Inches(0.42), Inches(0.68), Inches(10.8), Inches(0.36))
    r = sub.text_frame.paragraphs[0].add_run()
    _run(r, "For the people who run the warehouse — not a tech story.", 14, False, MUTED)

    _card(
        s,
        0.42,
        1.22,
        4.05,
        5.55,
        "Business case",
        [
            "We already paid for robots and warehouse software.",
            "When those tools disagree, people still send a robot.",
            "That costs time, late trucks, and unsafe moves.",
            "The case is simple: stop a bad send before it happens.",
        ],
        STEEL,
        15,
    )
    _card(
        s,
        4.62,
        1.22,
        4.05,
        5.55,
        "Problem statement",
        [
            "The warehouse already has robots.",
            "The screens often do not match the floor.",
            "People still press Go.",
        ],
        STOP,
        16,
    )
    _card(
        s,
        8.82,
        1.22,
        4.08,
        5.55,
        "Real-world problem we are solving",
        [
            "A supervisor must pick an order and make the truck.",
            "A robot looks free — but should stay still.",
            "A bin shows two or three different counts.",
            "An order looks gone. The truck did not leave.",
            "We solve this: one clear stop or go before anyone sends a robot.",
        ],
        AMBER,
        15,
    )
    s.notes_slide.notes_text_frame.text = (
        "Say: we already paid for robots. The screens lie. People still press Go. "
        "We want one stop or go before a robot moves."
    )


def _stat_box(slide, l, t, w, h, number, label, accent):
    _round(slide, l, t, w, h, CARD, LINE)
    _rect(slide, l, t, w, 0.07, accent)
    num = slide.shapes.add_textbox(Inches(l), Inches(t + 0.28), Inches(w), Inches(0.95))
    tf = num.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    _run(r, str(number), 36, True, accent)
    cap = slide.shapes.add_textbox(
        Inches(l + 0.10), Inches(t + 1.18), Inches(w - 0.20), Inches(h - 1.32)
    )
    tf = cap.text_frame
    tf.word_wrap = True
    for i, line in enumerate(label.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(0)
        r = p.add_run()
        _run(r, line, 13, False, INK)


def _add_fde_stats(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _chrome(s, "DockSight  ·  working deck  ·  FDE model vs this repo")

    title = s.shapes.add_textbox(Inches(2.72), Inches(0.16), Inches(8.8), Inches(0.42))
    r = title.text_frame.paragraphs[0].add_run()
    _run(r, "What we generated  ·  FDE model", 22, True, INK)

    sub = s.shapes.add_textbox(Inches(0.42), Inches(0.62), Inches(11.0), Inches(0.36))
    r = sub.text_frame.paragraphs[0].add_run()
    _run(
        r,
        "Compared to the 21-step FDE operating model. This is a virtual proof. Not live robots.",
        14,
        False,
        MUTED,
    )

    boxes = [
        ("21", "Steps in the\nFDE model", STEEL),
        ("16", "Steps done\nin this proof", GO),
        ("5", "Steps left for\nlive customer", AMBER),
        ("45", "Written\nartifacts", STEEL),
        ("8", "Stop rules\n(hard gates)", STOP),
        ("22", "Evals passed\n(0 fail)", GO),
        ("76", "Automated\ntests", STEEL),
        ("42", "Refuse-path\ntest cases", STOP),
        ("4", "Screens\nbuilt", AMBER),
        ("0", "Live robots\nmoved", MUTED),
    ]
    # 5 x 2 small boxes
    w, h, gap = 2.38, 2.42, 0.12
    x0, y0 = 0.42, 1.08
    for i, (num, label, accent) in enumerate(boxes):
        col = i % 5
        row = i // 5
        _stat_box(s, x0 + col * (w + gap), y0 + row * (h + 0.18), w, h, num, label, accent)

    s.notes_slide.notes_text_frame.text = (
        "21-step FDE PDF. We completed 1–16 for a virtual warehouse tower. "
        "17–21 (live deploy, value in dollars, ISO audit, retire) are not claimed. "
        "45 = discovery files + specs + two ADRs + traceability + component register. "
        "42 = named refuse cases in the negative test pack. "
        "76 = pytest functions. 22 evals all PASS. 4 screens: dashboard, workflow, browse, contention. "
        "Zero live robots: physical_control stays disabled."
    )


def _flow_step(slide, l, t, w, h, num_title, lines, file_name, accent):
    _round(slide, l, t, w, h, CARD, LINE)
    _rect(slide, l, t, w, 0.08, accent)
    head = slide.shapes.add_textbox(
        Inches(l + 0.08), Inches(t + 0.16), Inches(w - 0.16), Inches(0.70)
    )
    tf = head.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    _run(r, num_title, 14, True, accent)
    body = slide.shapes.add_textbox(
        Inches(l + 0.10), Inches(t + 0.88), Inches(w - 0.20), Inches(h - 1.55)
    )
    tf = body.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(4)
        r = p.add_run()
        _run(r, line, 12, False, INK)
    foot = slide.shapes.add_textbox(
        Inches(l + 0.08), Inches(t + h - 0.55), Inches(w - 0.16), Inches(0.48)
    )
    tf = foot.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    _run(r, file_name, 11, True, MUTED)


def _arrow(slide, l, t):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW, Inches(l), Inches(t), Inches(0.20), Inches(0.16)
    )
    _fill(shp, AMBER)


def _add_order_path(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _chrome(
        s,
        "DockSight  ·  working deck  ·  current path  ·  this app does not drive step 5",
    )

    title = s.shapes.add_textbox(Inches(2.72), Inches(0.16), Inches(8.8), Inches(0.42))
    r = title.text_frame.paragraphs[0].add_run()
    _run(r, "The path an order takes today", 22, True, INK)

    sub = s.shapes.add_textbox(Inches(0.42), Inches(0.60), Inches(11.0), Inches(0.32))
    r = sub.text_frame.paragraphs[0].add_run()
    _run(
        r,
        "Customer want  →  six hand-offs  →  truck left. This is the current walk, not a future design.",
        14,
        False,
        MUTED,
    )

    steps = [
        (
            "1  Order",
            ["OMS takes the order.", "Clock to the truck starts."],
            "orders.csv",
            STEEL,
        ),
        (
            "2  Allocate",
            ["WMS holds stock.", "ERP and camera also claim qty."],
            "inventory + skus",
            STEEL,
        ),
        (
            "3  Work",
            ["WES makes jobs.", "Pick, move, pack, stage."],
            "tasks.csv",
            STEEL,
        ),
        (
            "4  Robot",
            ["Fleet names a machine.", "Cert, repair, zone."],
            "robots + zones",
            STEEL,
        ),
        (
            "5  Move",
            ["Pick and conveyor.", "This screen does not drive it."],
            "telemetry / assets",
            STOP,
        ),
        (
            "6  Ship",
            ["TMS books dock + truck.", "Then the truck can leave."],
            "shipments.csv",
            GO,
        ),
    ]
    w, h = 1.88, 4.05
    x, y = 0.38, 0.98
    for i, (title_s, lines, fname, accent) in enumerate(steps):
        _flow_step(s, x, y, w, h, title_s, lines, fname, accent)
        if i < len(steps) - 1:
            _arrow(s, x + w + 0.01, y + 1.90)
        x += w + 0.22

    _round(s, 0.42, 5.18, 12.48, 1.76, CARD, LINE)
    bar = s.shapes.add_textbox(Inches(0.62), Inches(5.30), Inches(12.1), Inches(1.52))
    tf = bar.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    _run(r, "Start: the customer wants the order.", 14, True, STEEL)
    p = tf.add_paragraph()
    p.space_before = Pt(6)
    r = p.add_run()
    _run(
        r,
        "End: truck departed — actual_departure is filled. Empty actual means it has not left.",
        14,
        False,
        INK,
    )
    p = tf.add_paragraph()
    p.space_before = Pt(6)
    r = p.add_run()
    _run(
        r,
        "DockSight watches steps 1–4 and 6. Step 5 is the real floor. We do not move the robot from this screen.",
        14,
        False,
        STOP,
    )

    s.notes_slide.notes_text_frame.text = (
        "Title means current traversal, not the target architecture. "
        "Walk left to right. Pause on 5: we observe, we do not drive. "
        "If they ask files: orders, inventory/skus, tasks, robots/maintenance/zones, telemetry, shipments."
    )


def _mini(slide, l, t, w, h, title, sub, accent, fill=None):
    _round(slide, l, t, w, h, fill or CARD, LINE)
    _rect(slide, l, t, 0.07, h, accent)
    tb = slide.shapes.add_textbox(
        Inches(l + 0.14), Inches(t + 0.06), Inches(w - 0.22), Inches(h - 0.12)
    )
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    _run(r, title, 12, True, accent)
    if sub:
        for line in sub.split("\n"):
            p = tf.add_paragraph()
            p.space_before = Pt(2)
            r = p.add_run()
            _run(r, line, 11, False, INK)


def _add_business_arch(prs):
    """Same idea as the sample FMS picture: left systems, center product, right user, floor below."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _chrome(s, "DockSight  ·  working deck  ·  business architecture  ·  observe only")

    title = s.shapes.add_textbox(Inches(2.72), Inches(0.14), Inches(8.9), Inches(0.38))
    r = title.text_frame.paragraphs[0].add_run()
    _run(r, "Business architecture: who talks, who decides", 20, True, INK)

    sub = s.shapes.add_textbox(Inches(0.38), Inches(0.50), Inches(10.6), Inches(0.28))
    r = sub.text_frame.paragraphs[0].add_run()
    _run(
        r,
        "Same picture as a fleet stack — but the middle is a watch-and-stop tower, not a robot driver.",
        13,
        False,
        MUTED,
    )

    # Left — other systems
    _mini(s, 0.32, 0.84, 2.28, 0.42, "Other systems", "We do not own these", STEEL, AMBER_SOFT)
    left = [
        ("OMS", "takes the order"),
        ("WMS · ERP · camera", "three stock claims"),
        ("WES", "makes the jobs"),
        ("Fleet · CMMS", "machine + repair"),
        ("TMS", "dock and truck"),
        ("Emails / lists", "not official rules"),
    ]
    y = 1.32
    for name, blurb in left:
        _mini(s, 0.32, y, 2.28, 0.72, name, blurb, STEEL)
        y += 0.80

    # Arrows left → center
    for ay in (1.62, 3.22, 4.82):
        _arrow(s, 2.64, ay)

    # Center — our tower (like the sample's FMS box)
    _round(s, 2.92, 0.84, 7.08, 6.08, CARD, STEEL)
    _rect(s, 2.92, 0.84, 7.08, 0.40, STEEL)
    banner = s.shapes.add_textbox(Inches(3.02), Inches(0.88), Inches(6.88), Inches(0.34))
    p = banner.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    _run(r, "DockSight  ·  virtual control tower  ·  reads files only", 13, True, WHITE)

    inner = s.shapes.add_textbox(Inches(3.10), Inches(1.32), Inches(6.72), Inches(0.32))
    p = inner.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    _run(r, "Join on order · robot · task · sku+bin   —   keep both sides if they disagree", 12, False, MUTED)

    cells = [
        ("Identity", "Is this one robot?"),
        ("Stock", "Are the counts trusted?"),
        ("Jobs", "Do WES and fleet agree?"),
        ("Gates", "Cert · repair · zone"),
    ]
    cx = 3.08
    for name, blurb in cells:
        _mini(s, cx, 1.72, 1.64, 1.22, name, blurb, STEEL, AMBER_SOFT)
        cx += 1.72

    _round(s, 3.08, 3.08, 6.76, 1.28, AMBER_SOFT, AMBER)
    dec = s.shapes.add_textbox(Inches(3.22), Inches(3.18), Inches(6.48), Inches(1.08))
    tf = dec.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    _run(r, "Decision", 14, True, STEEL)
    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(4)
    r = p.add_run()
    _run(r, "ALLOW    ·    DENY    ·    wait", 16, True, INK)
    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(4)
    r = p.add_run()
    _run(r, "No chatbot on the send path. Assign stays blocked.", 12, False, MUTED)

    # Floor band — like sample "Mobile Robots" but we do not control them
    _rect(s, 2.92, 4.52, 7.08, 0.10, AMBER)
    _round(s, 3.08, 4.74, 6.76, 1.98, PAPER, LINE)
    fl = s.shapes.add_textbox(Inches(3.22), Inches(4.84), Inches(6.48), Inches(1.78))
    tf = fl.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    _run(r, "Physical floor  —  not driven from this screen", 13, True, STOP)
    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(6)
    r = p.add_run()
    _run(r, "Robots   ·   conveyor   ·   dock   ·   truck", 13, False, INK)
    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(6)
    r = p.add_run()
    _run(r, "Telemetry is a claim. We do not POST motion.", 12, False, MUTED)

    # Right — people (like sample User Interface)
    _mini(s, 10.18, 0.84, 2.72, 0.42, "People", "See stop or go", GO, AMBER_SOFT)
    _mini(s, 10.18, 1.36, 2.72, 1.15, "Supervisor", "The person who must press Go — or wait", STEEL)
    _mini(
        s,
        10.18,
        2.62,
        2.72,
        2.05,
        "Screens",
        "Home wall\nOrder path\nContention\nBrowse",
        STEEL,
    )
    _mini(s, 10.18, 4.78, 2.72, 0.95, "Files", "CSV + SQLite  ·  not cleaned", STEEL)
    _mini(s, 10.18, 5.84, 2.72, 1.08, "Out", "No live robot  ·  no WMS write", STOP)

    s.notes_slide.notes_text_frame.text = (
        "Mimic the sample: left = other systems, middle = product, right = user, bottom = machines. "
        "Difference: we are not a fleet OS. Middle only judges. Floor is disconnected for write. "
        "Say: OMS WMS WES fleet TMS already run. DockSight sits in the middle and says ALLOW, DENY, or wait."
    )


def _frame(slide, l, t, w, h, fill, line_c, dash=False, weight=1.5):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line_c
    shp.line.width = Pt(weight)
    if dash:
        shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    try:
        shp.adjustments[0] = 0.08
    except Exception:
        pass
    return shp


def _center_text(slide, l, t, w, h, text, size, bold, color):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    _run(r, text, size, bold, color)
    return tb


def _arch_light(slide, footer):
    """Clean floor — no racks or AMR so the diagram can breathe."""
    bg = _rect(slide, 0, 0, 13.333, 7.5, PAPER)
    _back(slide, bg)
    _rect(slide, 0, 0, 0.10, 7.5, STEEL)
    _rect(slide, 0, 7.12, 13.333, 0.10, AMBER)
    _rect(slide, 0, 7.22, 13.333, 0.28, STEEL_DEEP)
    _footer(slide, footer)


def _add_business_arch_v2(prs):
    """Clearer nested stack, same story as v1. Keep v1 in the deck."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _arch_light(s, "DockSight  ·  architecture v2  ·  claims in  ·  judge  ·  person sees  ·  floor not driven")

    _center_text(s, 0.35, 0.12, 12.6, 0.36, "Business architecture", 24, True, INK)
    _center_text(
        s,
        0.35,
        0.46,
        12.6,
        0.28,
        "Other systems claim   →   DockSight judges   →   supervisor sees stop or go",
        14,
        False,
        MUTED,
    )

    # LEFT cluster
    _frame(s, 0.28, 0.86, 2.42, 6.08, CARD, STEEL, weight=1.25)
    _rect(s, 0.28, 0.86, 2.42, 0.46, STEEL)
    _center_text(s, 0.28, 0.92, 2.42, 0.36, "Already running", 13, True, WHITE)

    left_rows = [
        ("OMS", "order"),
        ("WMS", "stock"),
        ("ERP + camera", "also claim qty"),
        ("WES", "jobs"),
        ("Fleet + CMMS", "machine + repair"),
        ("TMS", "dock + truck"),
        ("Emails", "not rules"),
    ]
    y = 1.44
    for name, blurb in left_rows:
        _round(s, 0.42, y, 2.14, 0.70, AMBER_SOFT, LINE)
        tb = s.shapes.add_textbox(Inches(0.54), Inches(y + 0.08), Inches(1.90), Inches(0.56))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        _run(r, name, 13, True, STEEL)
        p = tf.add_paragraph()
        r = p.add_run()
        _run(r, blurb, 11, False, INK)
        y += 0.76

    _arrow(s, 2.74, 3.70)

    # CENTER nested estate (sample: WES wraps FMS wraps robots)
    _frame(s, 3.02, 0.86, 7.12, 6.08, CARD, STEEL, dash=True, weight=1.75)
    _rect(s, 3.02, 0.86, 7.12, 0.38, AMBER_SOFT)
    _center_text(s, 3.02, 0.90, 7.12, 0.32, "Warehouse estate  ·  files in, nothing written back", 12, True, STEEL)

    _frame(s, 3.22, 1.36, 6.72, 3.18, WHITE, STEEL, weight=1.5)
    _rect(s, 3.22, 1.36, 6.72, 0.44, STEEL)
    _center_text(s, 3.22, 1.40, 6.72, 0.36, "DockSight  ·  watch-and-stop tower", 15, True, WHITE)

    tiles = [
        ("1  Who", "one robot?"),
        ("2  Stock", "counts agree?"),
        ("3  Jobs", "WES = fleet?"),
        ("4  Safe", "cert · zone"),
    ]
    tx = 3.38
    for head, blurb in tiles:
        _round(s, tx, 1.94, 1.52, 1.18, AMBER_SOFT, LINE)
        _center_text(s, tx, 2.04, 1.52, 0.42, head, 13, True, STEEL)
        _center_text(s, tx, 2.46, 1.52, 0.50, blurb, 12, False, INK)
        tx += 1.62

    _round(s, 3.38, 3.26, 6.40, 1.10, AMBER_SOFT, AMBER)
    _center_text(s, 3.38, 3.34, 6.40, 0.38, "Decision", 12, True, MUTED)
    _center_text(s, 3.38, 3.68, 6.40, 0.50, "ALLOW     DENY     wait", 18, True, INK)

    # Floor — sample's "Mobile Robots" band, but disconnected
    _rect(s, 3.22, 4.64, 6.72, 0.08, AMBER)
    _frame(s, 3.22, 4.82, 6.72, 1.92, PAPER, STOP, dash=True, weight=1.5)
    _center_text(s, 3.22, 4.90, 6.72, 0.32, "Floor  ·  we watch  ·  we do not drive", 13, True, STOP)
    floor = [("Robot", "AMR"), ("Conveyor", "plant"), ("Dock", "door"), ("Truck", "TMS")]
    fx = 3.42
    for name, blurb in floor:
        _round(s, fx, 5.30, 1.48, 1.22, CARD, LINE)
        _rect(s, fx, 5.30, 1.48, 0.08, AMBER)
        _center_text(s, fx, 5.46, 1.48, 0.42, name, 13, True, STEEL)
        _center_text(s, fx, 5.88, 1.48, 0.42, blurb, 11, False, MUTED)
        fx += 1.62

    # RIGHT user column (sample: UI → desktop/tablet → person)
    _frame(s, 10.32, 0.86, 2.62, 6.08, CARD, GO, weight=1.25)
    _rect(s, 10.32, 0.86, 2.62, 0.46, GO)
    _center_text(s, 10.32, 0.92, 2.62, 0.36, "People see", 13, True, WHITE)

    _round(s, 10.48, 1.48, 2.30, 1.55, AMBER_SOFT, LINE)
    _center_text(s, 10.48, 1.56, 2.30, 0.36, "Screens", 13, True, STEEL)
    _center_text(s, 10.48, 1.92, 2.30, 0.95, "wall  ·  path\ncontention  ·  browse", 12, False, INK)

    _round(s, 10.48, 3.18, 2.30, 1.35, CARD, LINE)
    _oval(s, 11.28, 3.28, 0.70, 0.70, STEEL)
    _center_text(s, 10.48, 4.02, 2.30, 0.40, "Supervisor", 13, True, STEEL)

    _round(s, 10.48, 4.68, 2.30, 0.95, CARD, LINE)
    _center_text(s, 10.48, 4.78, 2.30, 0.36, "Files", 12, True, STEEL)
    _center_text(s, 10.48, 5.12, 2.30, 0.40, "CSV  ·  SQLite", 12, False, INK)

    _round(s, 10.48, 5.78, 2.30, 0.96, AMBER_SOFT, STOP)
    _center_text(s, 10.48, 5.88, 2.30, 0.36, "Off", 12, True, STOP)
    _center_text(s, 10.48, 6.22, 2.30, 0.40, "no live send", 12, False, INK)

    _arrow(s, 10.08, 2.20)

    s.notes_slide.notes_text_frame.text = (
        "Better version: nested like the sample picture. Outer dashed = estate. "
        "Inner steel = DockSight. Amber bar = decision. Dashed red floor = machines not driven. "
        "Keep v1 if this feels too sparse. Present only one of the two."
    )


def _layer_row(slide, l, t, w, h, num, heading, body, accent, fill=None):
    _round(slide, l, t, w, h, fill or CARD, LINE)
    _rect(slide, l, t, 0.08, h, accent)
    tb = slide.shapes.add_textbox(
        Inches(l + 0.16), Inches(t + 0.06), Inches(w - 0.28), Inches(h - 0.12)
    )
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    _run(r, num + "  " + heading, 13, True, accent)
    p = tf.add_paragraph()
    p.space_before = Pt(2)
    r = p.add_run()
    _run(r, body, 12, False, INK)


def _add_sa_vs_fde(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _chrome(s, "DockSight  ·  working deck  ·  artifacts first  ·  AI last")

    title = s.shapes.add_textbox(Inches(2.72), Inches(0.16), Inches(8.8), Inches(0.42))
    r = title.text_frame.paragraphs[0].add_run()
    _run(r, "Same capstone, two ways to run it", 22, True, INK)

    sub = s.shapes.add_textbox(Inches(0.42), Inches(0.58), Inches(11.2), Inches(0.30))
    r = sub.text_frame.paragraphs[0].add_run()
    _run(
        r,
        "FDE does not jump to AI. It walks artifacts. AI is last, and it cannot send a robot.",
        14,
        False,
        MUTED,
    )

    _round(s, 0.38, 0.96, 6.18, 0.38, AMBER_SOFT, MUTED)
    _center_text(s, 0.38, 1.00, 6.18, 0.32, "Solution architect  ·  skip to the picture", 13, True, MUTED)

    left = [
        ("1", "Target drawing", "C4 and platform first."),
        ("2", "Trust the names", "WMS is stock. Fleet is free."),
        ("3", "Clean the files", "Make the KPI look green."),
        ("4", "Jump to AI", "Chatbot, copilot, agents."),
        ("5", "Hand off", "Someone else builds it later."),
    ]
    y = 1.42
    for num, head, body in left:
        _layer_row(s, 0.38, y, 6.18, 0.86, num, head, body, MUTED)
        y += 0.90

    _round(s, 6.76, 0.96, 6.16, 0.38, AMBER_SOFT, GO)
    _center_text(s, 6.76, 1.00, 6.16, 0.32, "FDE engineer  ·  artifacts, then a thin slice", 13, True, GO)

    right = [
        ("1", "Mandate + evidence", "Charter, RACI, field register. Named IDs."),
        ("2", "Process + RCA + KPIs", "SIPOC, conflicts, SCQA. No cleaned average."),
        ("3", "Domain + contracts", "Words and join keys. Keep both sides."),
        ("4", "Gates + evals + tests", "Stop rules before any helper. 22 evals."),
        ("5", "Working tower", "ALLOW / DENY / wait. Observe only."),
        ("6", "AI last — optional", "Explain a stop. Never write. Default off."),
    ]
    y = 1.42
    for num, head, body in right:
        _layer_row(s, 6.76, y, 6.16, 0.72, num, head, body, GO)
        y += 0.76

    _round(s, 0.38, 6.22, 12.54, 0.74, AMBER_SOFT, AMBER)
    _center_text(
        s,
        0.48,
        6.34,
        12.34,
        0.52,
        "Left skips to AI.  Right earns AI: files → named failures → gates → tests → then maybe a helper that cannot send.",
        14,
        True,
        STEEL,
    )

    s.notes_slide.notes_text_frame.text = (
        "FDE layers: mandate, evidence, process, RCA, domain, contracts, hard gates, evals, then a thin tower. "
        "Copilot is layer 6 and optional. It cannot assign. "
        "Architect path: drawing, trust names, clean data, chatbot. "
        "Say: we did not start with AI suggestions. We started with artifacts."
    )


def _add_current_kpis(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _chrome(s, "DockSight  ·  working deck  ·  Prompt 03 snapshot  ·  files not cleaned")

    title = s.shapes.add_textbox(Inches(2.72), Inches(0.16), Inches(8.8), Inches(0.42))
    r = title.text_frame.paragraphs[0].add_run()
    _run(r, "Current KPIs  ·  the inherited warehouse", 22, True, INK)

    sub = s.shapes.add_textbox(Inches(0.42), Inches(0.58), Inches(11.2), Inches(0.30))
    r = sub.text_frame.paragraphs[0].add_run()
    _run(
        r,
        "Same brownfield files. Not the 7,000-row overlay. After UC-1 we refuse — we do not rewrite these counts.",
        14,
        False,
        MUTED,
    )

    boxes = [
        ("28.5%", "Robot looks free\n203 / 712 should not move", STOP),
        ("78.9%", "Stock counts split\n7,382 / 9,360 bins", STOP),
        ("84.2%", "Job split\nWES ≠ fleet  7,351", STOP),
        ("36.7%", "Order split\nOMS ≠ WMS  1,283", AMBER),
        ("16.7%", "Truck already late\n586 DELAYED", AMBER),
        ("0", "Orders that close\norder → truck clean", STEEL),
        ("35", "Expired cert\nstill connected", STOP),
        ("176", "Open repair\nstill marked free", STOP),
        ("0", "Unsafe assigns\non the new path", GO),
        ("—", "Use / recovery\nnot computable", MUTED),
    ]
    w, h, gap = 2.38, 2.42, 0.12
    x0, y0 = 0.42, 0.98
    for i, (num, label, accent) in enumerate(boxes):
        col = i % 5
        row = i // 5
        _stat_box(s, x0 + col * (w + gap), y0 + row * (h + 0.16), w, h, num, label, accent)

    s.notes_slide.notes_text_frame.text = (
        "Do not put 84% as-of cutoff on this slide. That is PARTIAL, not a board KPI. "
        "Do not use teammate 62% on-time. "
        "After is treatment: 203 still in the file, 0 treated as eligible. "
        "Unsafe assign 0 is the CTQ on UC-1, not a claim the estate is clean."
    )


def _add_failure_modes(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _chrome(s, "DockSight  ·  working deck  ·  failure modes  ·  synthetic fixtures")

    title = s.shapes.add_textbox(Inches(2.72), Inches(0.16), Inches(8.8), Inches(0.42))
    r = title.text_frame.paragraphs[0].add_run()
    _run(r, "Failure modes — and how FDE stops them", 22, True, INK)

    sub = s.shapes.add_textbox(Inches(0.42), Inches(0.58), Inches(11.2), Inches(0.30))
    r = sub.text_frame.paragraphs[0].add_run()
    _run(
        r,
        "Named rows from the brownfield files and the observe-only dashboard. We do not clean them. We refuse.",
        14,
        False,
        MUTED,
    )

    # Header row
    _round(s, 0.38, 0.98, 4.15, 0.40, STOP, STOP)
    _center_text(s, 0.38, 1.02, 4.15, 0.32, "What breaks on the floor", 13, True, WHITE)
    _round(s, 4.68, 0.98, 3.70, 0.40, STEEL, STEEL)
    _center_text(s, 4.68, 1.02, 3.70, 0.32, "What the data shows", 13, True, WHITE)
    _round(s, 8.52, 0.98, 4.40, 0.40, GO, GO)
    _center_text(s, 8.52, 1.02, 4.40, 0.32, "What FDE does", 13, True, WHITE)

    rows = [
        (
            "Robot looks free — should stay still",
            "RBT-0020 EXPIRED\nRBT-0001 open repair",
            "G1 / G2 DENY\nAssign blocked on wall",
        ),
        (
            "Bin shows three counts",
            "SKU-01146\n205 / 205 / 202",
            "Mark UNCERTAIN\nNo pick-as-known",
        ),
        (
            "Order looks shipped — job still runs",
            "ORD-000968 SHIPPED\nTSK still EXECUTING",
            "Not complete (G6)\nKeep both statuses",
        ),
        (
            "Truck already late — screens look fine",
            "ORD-000004\nTMS DELAYED, no actual",
            "on_time = false\nMiss cutoff, not zone release",
        ),
        (
            "Two machines, one radio name",
            "BOT-COLLISION-01\non RBT-0001 and RBT-0002",
            "Keep both IDs\nNever merge or guess",
        ),
        (
            "Path into a closed aisle",
            "TSK → DC-01-Z05\nRESTRICTED",
            "G4 refuse path\nRelease zone = DENY",
        ),
    ]
    y = 1.48
    for fail, data, fix in rows:
        _round(s, 0.38, y, 4.15, 0.78, CARD, LINE)
        _rect(s, 0.38, y, 0.08, 0.78, STOP)
        tb = s.shapes.add_textbox(Inches(0.56), Inches(y + 0.12), Inches(3.85), Inches(0.58))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        _run(r, fail, 13, True, INK)

        _round(s, 4.68, y, 3.70, 0.78, CARD, LINE)
        tb = s.shapes.add_textbox(Inches(4.82), Inches(y + 0.08), Inches(3.42), Inches(0.64))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, line in enumerate(data.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_before = Pt(1)
            r = p.add_run()
            _run(r, line, 12, False, STEEL)

        _round(s, 8.52, y, 4.40, 0.78, AMBER_SOFT, LINE)
        tb = s.shapes.add_textbox(Inches(8.66), Inches(y + 0.08), Inches(4.12), Inches(0.64))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, line in enumerate(fix.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_before = Pt(1)
            r = p.add_run()
            _run(r, line, 12, True if i == 0 else False, GO if i == 0 else INK)

        y += 0.84

    _round(s, 0.38, 6.58, 12.54, 0.42, AMBER_SOFT, AMBER)
    _center_text(
        s,
        0.48,
        6.64,
        12.34,
        0.32,
        "Dashboard shows the mess. Green control = ORD-009999. Broken path = ORD-000004. We do not drive the robot.",
        13,
        True,
        STEEL,
    )

    s.notes_slide.notes_text_frame.text = (
        "Walk three columns: failure, named fixture, FDE refusal. "
        "Demo: open dashboard or /workflow — ORD-009999 then ORD-000004. "
        "Overcome does not mean fix the CSV. It means stop, keep both sides, evals pass."
    )


def _add_failure_sell(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _chrome(s, "DockSight  ·  working deck  ·  stakeholder  ·  process addresses the failure")

    title = s.shapes.add_textbox(Inches(2.72), Inches(0.16), Inches(8.8), Inches(0.42))
    r = title.text_frame.paragraphs[0].add_run()
    _run(r, "This is what goes wrong — and what we put in front of it", 20, True, INK)

    sub = s.shapes.add_textbox(Inches(0.42), Inches(0.56), Inches(11.2), Inches(0.30))
    r = sub.text_frame.paragraphs[0].add_run()
    _run(
        r,
        "For the people who own the warehouse. Not a tech list. The process is: see it, name it, stop the bad send.",
        14,
        False,
        MUTED,
    )

    _round(s, 0.38, 0.94, 6.18, 0.38, STOP, STOP)
    _center_text(s, 0.38, 0.98, 6.18, 0.30, "If we do nothing", 13, True, WHITE)
    _round(s, 6.76, 0.94, 6.16, 0.38, GO, GO)
    _center_text(s, 6.76, 0.98, 6.16, 0.30, "What our process does", 13, True, WHITE)

    left = [
        ("Wrong robot still goes", "Looks free. Cert expired or still in repair."),
        ("Pick that is not there", "Three counts on one bin. Someone still picks."),
        ("Customer thinks it left", "Order looks shipped. Truck did not leave."),
        ("Late truck, no early warning", "Cutoff already missed. Screens still look fine."),
        ("People paper over it", "Emails and unofficial lists become the real system."),
    ]
    y = 1.40
    for head, body in left:
        _round(s, 0.38, y, 6.18, 0.90, CARD, LINE)
        _rect(s, 0.38, y, 0.08, 0.90, STOP)
        tb = s.shapes.add_textbox(Inches(0.58), Inches(y + 0.10), Inches(5.84), Inches(0.72))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        _run(r, head, 14, True, STOP)
        p = tf.add_paragraph()
        p.space_before = Pt(3)
        r = p.add_run()
        _run(r, body, 13, False, INK)
        y += 0.96

    right = [
        ("Stop before Go", "One path: order → stock → job → robot → truck. If they disagree, wait."),
        ("Named proof, not a slogan", "We show ORD-000004 broken and ORD-009999 agreed — on the same screen."),
        ("Rules first", "ALLOW, DENY, or wait. No chatbot that sends a robot."),
        ("Keep the conflict visible", "We do not pick a winner system or clean the books for a nicer KPI."),
        ("Ask of you", "Accept this as the control idea. Do not turn robots on for the showcase."),
    ]
    y = 1.40
    for head, body in right:
        _round(s, 6.76, y, 6.16, 0.90, CARD, LINE)
        _rect(s, 6.76, y, 0.08, 0.90, GO)
        tb = s.shapes.add_textbox(Inches(6.96), Inches(y + 0.10), Inches(5.82), Inches(0.72))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        _run(r, head, 14, True, GO)
        p = tf.add_paragraph()
        p.space_before = Pt(3)
        r = p.add_run()
        _run(r, body, 13, False, INK)
        y += 0.96

    s.notes_slide.notes_text_frame.text = (
        "Stakeholder / mimicking-client slide. Sell the failure, then the process. "
        "Do not say G1 G2 UNCERTAIN. Say: wrong robot, fake stock, fake shipped, late truck, shadow emails. "
        "Process: see, name, stop. Demo ORD-000004 vs ORD-009999 if they ask to see it. "
        "Ask: accept the control idea. Robots stay off."
    )


def _add_title_selection(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _chrome(s, "DockSight  ·  FDE capstone  ·  synthetic proof  ·  not live robots")

    title = s.shapes.add_textbox(Inches(2.72), Inches(0.16), Inches(8.8), Inches(0.42))
    r = title.text_frame.paragraphs[0].add_run()
    _run(r, "DockSight  ·  why this capstone", 22, True, INK)

    sub = s.shapes.add_textbox(Inches(0.42), Inches(0.58), Inches(11.2), Inches(0.32))
    r = sub.text_frame.paragraphs[0].add_run()
    _run(
        r,
        "We chose a brownfield warehouse with robots already in it — not a greenfield AI demo.",
        14,
        False,
        MUTED,
    )

    _card(
        s,
        0.42,
        1.08,
        4.05,
        5.70,
        "Why this problem",
        [
            "The estate already paid for robots and many systems.",
            "Exceptions — not the happy path — burn people and create late trucks and unsafe sends.",
            "A supervisor still has no one screen that says stop or go.",
            "That is a real customer problem, not a missing chatbot.",
        ],
        STEEL,
        14,
    )
    _card(
        s,
        4.62,
        1.08,
        4.05,
        5.70,
        "What we selected",
        [
            "A virtual watch-and-stop tower.",
            "Rules first (Option A).",
            "Named journeys: ORD-000004 broken, ORD-009999 agreed.",
            "Evals and tests before anything that looks autonomous.",
            "Observe only. Physical control stays off.",
        ],
        GO,
        14,
    )
    _card(
        s,
        8.82,
        1.08,
        4.08,
        5.70,
        "What we refused to pick",
        [
            "An agent that dispatches robots.",
            "A mandatory graph, twin, or RAG.",
            "Cleaning CSVs so KPIs look green.",
            "Turning robots on for a nicer demo.",
            "Dollar ROI we cannot compute.",
        ],
        STOP,
        14,
    )
    s.notes_slide.notes_text_frame.text = (
        "Capstone selection. Say: we did not pick the easiest AI story. We picked the inherited mess. "
        "Selection is Option A deterministic tower. Kill list is as important as the build."
    )


def _add_trust_autonomy(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _chrome(s, "DockSight  ·  trust  ·  autonomy  ·  not production")

    title = s.shapes.add_textbox(Inches(2.72), Inches(0.16), Inches(8.8), Inches(0.42))
    r = title.text_frame.paragraphs[0].add_run()
    _run(r, "Trust, autonomy, and production honesty", 22, True, INK)

    sub = s.shapes.add_textbox(Inches(0.42), Inches(0.58), Inches(11.2), Inches(0.30))
    r = sub.text_frame.paragraphs[0].add_run()
    _run(
        r,
        "Trainers score this as: we know what software may do, and we did not fake a live plant.",
        14,
        False,
        MUTED,
    )

    rows = [
        ("Watch", "Read the files. Show conflicts.", "On now"),
        ("Advise", "Miss cutoff. Do not assign this robot.", "On now — human may ignore"),
        ("Preview only", "ALLOW on screen. Nothing is applied.", "On now"),
        ("Safety / motion", "E-stop, zone release, send a robot.", "Software DENY. Always."),
    ]
    y = 1.02
    _round(s, 0.42, y, 3.90, 0.40, STEEL, STEEL)
    _center_text(s, 0.42, y + 0.04, 3.90, 0.32, "What software may do", 13, True, WHITE)
    _round(s, 4.46, y, 5.20, 0.40, STEEL, STEEL)
    _center_text(s, 4.46, y + 0.04, 5.20, 0.32, "Meaning", 13, True, WHITE)
    _round(s, 9.80, y, 3.10, 0.40, STEEL, STEEL)
    _center_text(s, 9.80, y + 0.04, 3.10, 0.32, "In this proof", 13, True, WHITE)

    y = 1.52
    colors = [GO, GO, AMBER, STOP]
    for (head, body, now), col in zip(rows, colors):
        _round(s, 0.42, y, 3.90, 0.78, CARD, LINE)
        _rect(s, 0.42, y, 0.08, 0.78, col)
        tb = s.shapes.add_textbox(Inches(0.62), Inches(y + 0.22), Inches(3.55), Inches(0.40))
        r = tb.text_frame.paragraphs[0].add_run()
        _run(r, head, 15, True, col)
        _round(s, 4.46, y, 5.20, 0.78, CARD, LINE)
        tb = s.shapes.add_textbox(Inches(4.62), Inches(y + 0.22), Inches(4.90), Inches(0.40))
        r = tb.text_frame.paragraphs[0].add_run()
        _run(r, body, 14, False, INK)
        _round(s, 9.80, y, 3.10, 0.78, AMBER_SOFT if col != STOP else CARD, LINE)
        tb = s.shapes.add_textbox(Inches(9.92), Inches(y + 0.22), Inches(2.86), Inches(0.40))
        r = tb.text_frame.paragraphs[0].add_run()
        _run(r, now, 13, True, col)
        y += 0.86

    _round(s, 0.42, 5.08, 12.48, 1.88, CARD, LINE)
    tb = s.shapes.add_textbox(Inches(0.62), Inches(5.20), Inches(12.10), Inches(1.64))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    _run(r, "Trust rules we will say out loud", 14, True, STEEL)
    for line in (
        "No system is true because of its name. LLM is off the write path. Email is not an approval.",
        "22 evals passed. Assign / pick / zone release stay blocked. physical_control stays disabled.",
        "This is not production. Production needs a new yes, a safety case, and live data — not a demo switch.",
    ):
        p = tf.add_paragraph()
        p.space_before = Pt(6)
        r = p.add_run()
        _run(r, line, 14, False, INK)

    s.notes_slide.notes_text_frame.text = (
        "Autonomy: observe and recommend only. No autonomous assign. "
        "Trust: gates, evals, no silent merge. Production: deferred 17–21. Do not claim ISO or live OT."
    )


def _add_moonshot(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _chrome(s, "DockSight  ·  moonshot is later  ·  not in this proof")

    title = s.shapes.add_textbox(Inches(2.72), Inches(0.16), Inches(8.8), Inches(0.42))
    r = title.text_frame.paragraphs[0].add_run()
    _run(r, "Innovation now  ·  moonshot next", 22, True, INK)

    sub = s.shapes.add_textbox(Inches(0.42), Inches(0.58), Inches(11.2), Inches(0.30))
    r = sub.text_frame.paragraphs[0].add_run()
    _run(
        r,
        "The brave idea is not “add AI.” It is: refuse to send until the story is true enough.",
        14,
        False,
        MUTED,
    )

    _card(
        s,
        0.42,
        1.08,
        6.15,
        5.70,
        "Innovation in this capstone (earned)",
        [
            "Treat disagreement as a first-class fact — not a bug to hide.",
            "Stop is a valid product answer.",
            "Filter then score — do not let “highest battery” win.",
            "Identity collisions stay visible. We do not guess AMR-044.",
            "A helper may explain a stop later. It still cannot send.",
        ],
        GO,
        15,
    )
    _card(
        s,
        6.77,
        1.08,
        6.12,
        5.70,
        "Moonshot — only after the tower holds",
        [
            "Explain-only copilot on a refuse (never on assign).",
            "Demand spike: re-rank eligible robots only — still no OT.",
            "Learn which conflicts keep coming back (still keep both qtys).",
            "Not this room: self-healing plant, agent mesh, live dispatch.",
            "If the model is down, gates still run. That is the design.",
        ],
        AMBER,
        15,
    )
    s.notes_slide.notes_text_frame.text = (
        "Moonshot scoring: show we thought beyond the demo without claiming we built it. "
        "Innovation is abstain + artifacts. AI is layer last."
    )


def _add_close_ask(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _chrome(s, "DockSight  ·  ask  ·  90 days only if you say so")

    title = s.shapes.add_textbox(Inches(2.72), Inches(0.16), Inches(8.8), Inches(0.42))
    r = title.text_frame.paragraphs[0].add_run()
    _run(r, "The story, and what we need from you", 22, True, INK)

    sub = s.shapes.add_textbox(Inches(0.42), Inches(0.58), Inches(11.2), Inches(0.30))
    r = sub.text_frame.paragraphs[0].add_run()
    _run(
        r,
        "Evidence first. Then a stop. Then a decision — not a live fleet.",
        14,
        False,
        MUTED,
    )

    _card(
        s,
        0.42,
        1.08,
        6.15,
        3.55,
        "Data story in one breath",
        [
            "203 of 712 robots look free and should not move.",
            "7,382 bins disagree. 586 trucks already late.",
            "Zero orders close the chain cleanly in the inherited files.",
            "We still show one agreed path: ORD-009999.",
            "We still show one broken path: ORD-000004. Assign stays blocked.",
        ],
        STEEL,
        14,
    )
    _card(
        s,
        6.77,
        1.08,
        6.12,
        3.55,
        "Ask of the stakeholder",
        [
            "Accept: this is the problem and a safe control idea.",
            "Do not turn robots on for the showcase.",
            "Start a 90-day track only if you say so — still synthetic at first.",
            "We will not claim production, ISO, or dollar ROI today.",
        ],
        GO,
        14,
    )
    _card(
        s,
        0.42,
        4.78,
        12.48,
        2.08,
        "If they only remember three things",
        [
            "Screens and the floor often do not agree — people still press Go.",
            "DockSight is a watch-and-stop tower, not a robot driver.",
            "FDE started with artifacts and named failures, not with an AI suggestion.",
        ],
        AMBER,
        15,
    )
    s.notes_slide.notes_text_frame.text = (
        "Close. Point to live demo if time. Ask for acceptance of the control idea. "
        "90-day clock starts only after they say yes."
    )


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    _add_title_selection(prs)
    _add_problem_horizontal(prs)
    _add_failure_sell(prs)
    _add_order_path(prs)
    _add_business_arch_v2(prs)
    _add_current_kpis(prs)
    _add_failure_modes(prs)
    _add_sa_vs_fde(prs)
    _add_fde_stats(prs)
    _add_trust_autonomy(prs)
    _add_moonshot(prs)
    _add_close_ask(prs)
    _add_business_arch(prs)
    _add_theme_sample(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
