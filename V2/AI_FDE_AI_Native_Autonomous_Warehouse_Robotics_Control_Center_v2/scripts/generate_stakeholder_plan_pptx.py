"""Two-slide stakeholder deck: current problem + production-ready timeline."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

OUT = (
    Path(__file__).resolve().parents[1]
    / "docs"
    / "STAKEHOLDER_Problem_and_Production_Plan.pptx"
)

BG = RGBColor(0x12, 0x16, 0x1C)
PANEL = RGBColor(0x1B, 0x22, 0x2B)
LINE = RGBColor(0x2C, 0x36, 0x42)
TEXT = RGBColor(0xE9, 0xEE, 0xF3)
MUTED = RGBColor(0x8D, 0x9A, 0xAA)
ACCENT = RGBColor(0x4A, 0x7F, 0xA8)
WARN = RGBColor(0xC9, 0x92, 0x2A)
DENY = RGBColor(0xC4, 0x4C, 0x44)
OK = RGBColor(0x3C, 0x8F, 0x68)


def _set_run(run, text, size=16, bold=False, color=TEXT):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _slide_bg(slide, footer):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    _fill(bg, BG)
    spTree = slide.shapes._spTree
    sp = bg._element
    spTree.remove(sp)
    spTree.insert(2, sp)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), Inches(7.5))
    _fill(bar, ACCENT)
    foot = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Inches(7.22), Inches(13.333), Inches(0.28)
    )
    _fill(foot, PANEL)
    tb = slide.shapes.add_textbox(Inches(0.4), Inches(7.24), Inches(12.5), Inches(0.22))
    r = tb.text_frame.paragraphs[0].add_run()
    _set_run(r, footer, 10, False, MUTED)


def _title(slide, text):
    box = slide.shapes.add_textbox(Inches(0.45), Inches(0.22), Inches(12.4), Inches(0.5))
    r = box.text_frame.paragraphs[0].add_run()
    _set_run(r, text, 26, True, TEXT)


def _subtitle(slide, text):
    box = slide.shapes.add_textbox(Inches(0.45), Inches(0.7), Inches(12.4), Inches(0.42))
    tf = box.text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    _set_run(r, text, 15, False, MUTED)


def _card(slide, left, top, w, h, heading, lines, heading_color=ACCENT):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(h)
    )
    _fill(shp, PANEL)
    shp.line.color.rgb = LINE
    tb = slide.shapes.add_textbox(
        Inches(left + 0.16), Inches(top + 0.12), Inches(w - 0.32), Inches(h - 0.22)
    )
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    _set_run(r, heading, 14, True, heading_color)
    for line in lines:
        p = tf.add_paragraph()
        p.space_before = Pt(5)
        r = p.add_run()
        _set_run(r, line, 13, False, TEXT)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # --- Slide 1: current problem ---
    s = prs.slides.add_slide(blank)
    _slide_bg(
        s,
        "Stakeholder briefing  ·  Synthetic proof  ·  Not live robots  ·  Not production",
    )
    _title(s, "The problem today")
    _subtitle(
        s,
        "The warehouse already has robots and many computer systems. They often do not agree — and people still hit Go.",
    )

    _card(
        s,
        0.4,
        1.22,
        6.2,
        2.35,
        "What goes wrong",
        [
            "A robot looks free but should not move (expired safety check or open repair).",
            "The same product in the same bin shows three different quantities.",
            "An order looks shipped or on time while the robot task is still running, or the truck is already late.",
            "Two machines can share the same nickname on the floor radio.",
        ],
        DENY,
    )
    _card(
        s,
        6.8,
        1.22,
        6.1,
        2.35,
        "Why it matters",
        [
            "Wrong robot sent to a job.",
            "Pick that is not really there.",
            "Customer thinks the order left when it did not.",
            "People paper over this with emails and unofficial lists.",
        ],
        WARN,
    )
    _card(
        s,
        0.4,
        3.72,
        12.5,
        3.25,
        "What we built in this proof (not production)",
        [
            "A watch-and-stop tower: if the data is not trusted, we do not assign, we do not pick as known, we do not open a blocked area, and we do not move a live robot.",
            "One clean example order (ORD-009999) shows what “agreed” looks like. The messy rows stay messy on purpose — we did not clean the books.",
            "We are not making robots faster. We are stopping a bad go-ahead when digital state is not physical truth.",
            "Ask of you today: this shows the problem and a safe control idea. The warehouse is not modernized yet.",
        ],
        OK,
    )

    # --- Slide 2: production-ready timeline ---
    s = prs.slides.add_slide(blank)
    _slide_bg(
        s,
        "Clock starts after you accept this proof  ·  Robots stay off until a separate safety case",
    )
    _title(s, "Path to production-ready")
    _subtitle(
        s,
        "After this demo. No AI that dispatches. No live robot control in the showcase.",
    )

    _card(
        s,
        0.4,
        1.18,
        4.05,
        4.55,
        "Days 1–30  ·  Same proof, real operator path",
        [
            "Stop using “highest battery wins” as the picker people see.",
            "Always show who the robot really is. If two names collide, keep both.",
            "Approvals are a real record — not an email that says “I approved it.”",
            "Put honest KPIs on the board. Do not use a fake 84% on-time number.",
            "Still synthetic. Still no live fleet.",
        ],
        ACCENT,
    )
    _card(
        s,
        4.6,
        1.18,
        4.05,
        4.55,
        "Days 31–60  ·  How people work, still no robots",
        [
            "Optional helper that explains a stop in plain language. It cannot send a robot.",
            "When quantities disagree: people count the bin. Do not secretly overwrite the warehouse system.",
            "Agree clocks and timezones for a future live feed.",
            "Agree mission IDs so we never replay a whole queue after an outage.",
        ],
        WARN,
    )
    _card(
        s,
        8.8,
        1.18,
        4.1,
        4.55,
        "Days 61–90  ·  Only with a new mandate",
        [
            "Train exception staff on this tower.",
            "Measure on a real slice of live data (new consent).",
            "Only then discuss whether robot control could ever turn on — with a safety case.",
            "Until then: watch, refuse, recommend miss the cutoff.",
            "We will not guess mystery robot names or pretend all quantities match.",
        ],
        OK,
    )

    note = s.shapes.add_textbox(Inches(0.45), Inches(5.85), Inches(12.4), Inches(1.2))
    tf = note.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    _set_run(r, "Decision we need from you", 13, True, TEXT)
    p = tf.add_paragraph()
    p.space_before = Pt(6)
    r = p.add_run()
    _set_run(
        r,
        "This proof is enough to show the problem and a safe control idea.  ·  The warehouse is not production-ready.  ·  Do not turn on robots for the showcase.  ·  Start the 90-day track only if you say so.",
        13,
        False,
        MUTED,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(path)
