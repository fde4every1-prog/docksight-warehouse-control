"""Generate a downloadable PRD PowerPoint from Repo 2 (virtual product only)."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parents[1] / "docs" / "PRD_Virtual_Warehouse_Control_Tower.pptx"

BG = RGBColor(0x12, 0x16, 0x1C)
PANEL = RGBColor(0x1B, 0x22, 0x2B)
LINE = RGBColor(0x2C, 0x36, 0x42)
TEXT = RGBColor(0xE9, 0xEE, 0xF3)
MUTED = RGBColor(0x8D, 0x9A, 0xAA)
ACCENT = RGBColor(0x4A, 0x7F, 0xA8)
WARN = RGBColor(0xC9, 0x92, 0x2A)
DENY = RGBColor(0xC4, 0x4C, 0x44)
OK = RGBColor(0x3C, 0x8F, 0x68)


def _set_run(run, text, size=16, bold=False, color=TEXT):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _slide_bg(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    _fill(bg, BG)
    spTree = slide.shapes._spTree
    sp = bg._element
    spTree.remove(sp)
    spTree.insert(2, sp)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), Inches(7.5))
    _fill(bar, ACCENT)
    foot = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Inches(7.22), Inches(13.333), Inches(0.28)
    )
    _fill(foot, PANEL)
    tb = slide.shapes.add_textbox(Inches(0.4), Inches(7.24), Inches(12.5), Inches(0.22))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    _set_run(
        r,
        "PRD  ·  Virtual warehouse control tower  ·  Synthetic / physical_control disabled  ·  Not live OT",
        10,
        False,
        MUTED,
    )


def _title(slide, text, y=0.28):
    box = slide.shapes.add_textbox(Inches(0.45), Inches(y), Inches(12.4), Inches(0.55))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    _set_run(r, text, 26, True, TEXT)
    return box


def _subtitle(slide, text, y=0.78):
    box = slide.shapes.add_textbox(Inches(0.45), Inches(y), Inches(12.4), Inches(0.4))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    _set_run(r, text, 14, False, MUTED)


def _bullets(slide, items, left=0.45, top=1.25, width=12.4, height=5.6, size=16):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(8)
        r = p.add_run()
        _set_run(r, item, size, False, TEXT)


def _card(slide, left, top, w, h, heading, lines, heading_color=ACCENT):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(h)
    )
    _fill(shp, PANEL)
    shp.line.color.rgb = LINE
    tb = slide.shapes.add_textbox(
        Inches(left + 0.16), Inches(top + 0.12), Inches(w - 0.32), Inches(h - 0.22)
    )
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    _set_run(r, heading, 14, True, heading_color)
    for line in lines:
        p = tf.add_paragraph()
        p.space_before = Pt(6)
        r = p.add_run()
        _set_run(r, line, 13, False, TEXT)


def _table(slide, left, top, width, rows, col_w=None):
    cols = len(rows[0])
    n = len(rows)
    table_shape = slide.shapes.add_table(n, cols, Inches(left), Inches(top), Inches(width), Inches(0.38 * n))
    table = table_shape.table
    if col_w:
        for i, w in enumerate(col_w):
            table.columns[i].width = Inches(w)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            r = p.add_run()
            _set_run(r, val, 11, ri == 0, TEXT if ri else TEXT)
            if ri == 0:
                r.font.bold = True
                r.font.color.rgb = TEXT
            fill = RGBColor(0x24, 0x2E, 0x3A) if ri == 0 else (PANEL if ri % 2 else RGBColor(0x16, 0x1C, 0x24))
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
    return table_shape


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 title
    s = prs.slides.add_slide(blank)
    _slide_bg(s)
    box = s.shapes.add_textbox(Inches(0.55), Inches(2.0), Inches(12.2), Inches(1.2))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    _set_run(r, "Product Requirements Document", 36, True, TEXT)
    box = s.shapes.add_textbox(Inches(0.55), Inches(3.2), Inches(12.2), Inches(1.4))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    _set_run(r, "Virtual warehouse robotics control tower", 22, False, ACCENT)
    p = tf.add_paragraph()
    r = p.add_run()
    _set_run(r, "Repo 2  ·  Deterministic observe / refuse  ·  UC-1", 16, False, MUTED)
    p = tf.add_paragraph()
    r = p.add_run()
    _set_run(r, "Synthetic data only. No live robots, warehouse, or OT.", 16, False, WARN)

    # 2 one sentence
    s = prs.slides.add_slide(blank)
    _slide_bg(s)
    _title(s, "1. Product in one sentence")
    _subtitle(s, "What we are building — and what we are not")
    _card(
        s,
        0.45,
        1.35,
        12.4,
        2.2,
        "Product",
        [
            "A virtual, read-only control tower that observes conflicts in brownfield warehouse data",
            "and refuses unsafe work: expired-cert robots, invented inventory qty, false order complete,",
            "and safety-zone bypass to hit cutoff. Physical control stays disabled.",
        ],
        OK,
    )
    _card(
        s,
        0.45,
        3.75,
        6.0,
        2.9,
        "Selected (ADR-001)",
        [
            "Option A — deterministic eligibility,",
            "inventory uncertainty, filter-then-score.",
            "LLM off the write path (ADR-002).",
            "Optional explain-only copilot: default OFF.",
        ],
        ACCENT,
    )
    _card(
        s,
        6.7,
        3.75,
        6.15,
        2.9,
        "Killed",
        [
            "Agent that dispatches robots.",
            "Mandatory knowledge graph / twin / RAG.",
            "Enabling physical_control for a demo.",
            "Cleaning CSVs to make KPIs look green.",
        ],
        DENY,
    )

    # 3 problem
    s = prs.slides.add_slide(blank)
    _slide_bg(s)
    _title(s, "2. Problem")
    _subtitle(s, "Inherited virtual estate — digital state ≠ physical state; available ≠ suitable")
    _table(
        s,
        0.45,
        1.3,
        12.4,
        [
            ["KPI (Prompt 03 formulas)", "Snapshot", "What the old path does"],
            ["False-available robots (expired or open CMMS)", "203 / 712 (28.5%)", "legacy_score can still pick them"],
            ["Inventory WMS ≠ ERP ≠ vision", "7,382 / 9,360 (78.9%)", "Treats WMS as the only qty"],
            ["WES status ≠ fleet status", "7,351 / 8,732 (84.2%)", "Can look shipped while pack executes"],
            ["Shipments DELAYED (TMS)", "586 / 3,500 (16.7%)", "OMS SLA can look fine"],
        ],
        [4.2, 3.6, 4.6],
    )
    _bullets(
        s,
        [
            "Challenge brief: exceptions consume disproportionate effort and create service, safety, and resilience risk.",
            "Do not use the 84% as-of cutoff pile as a board KPI (PARTIAL historical artifact).",
        ],
        top=5.55,
        height=1.4,
        size=14,
    )

    # 4 goals
    s = prs.slides.add_slide(blank)
    _slide_bg(s)
    _title(s, "3. Goals and non-goals")
    _card(
        s,
        0.45,
        1.25,
        6.1,
        5.5,
        "Goals (Repo 2 / virtual product)",
        [
            "Refuse expired-cert and open-CMMS assigns on the new path (unsafe assign = 0 in fixtures).",
            "Show disagreeing inventory as UNCERTAIN; never emit one invented qty.",
            "Do not close ORD-000968 while pack-feed is EXECUTING.",
            "Recommend miss-cutoff; DENY safety-zone release.",
            "Evals (22/22) before any action that looks autonomous.",
            "Keep physical_control disabled. Keep brownfield CSVs unchanged.",
        ],
        OK,
    )
    _card(
        s,
        6.75,
        1.25,
        6.1,
        5.5,
        "Non-goals",
        [
            "Live robots, PLC, WMS write, fleet POST /missions.",
            "Dollar ROI / headcount reduction (no cost field in snapshot).",
            "Making DELAYED 586 → 0 in the CSV.",
            "ISO/IEC 42001 certification.",
            "Repo 3 / optional copilot unless a new mandate.",
            "Guessing AMR-044 = RBT-0044.",
        ],
        DENY,
    )

    # 5 users
    s = prs.slides.add_slide(blank)
    _slide_bg(s)
    _title(s, "4. Users and journeys")
    _subtitle(s, "Virtual roles only — no connected floor")
    _card(s, 0.45, 1.25, 4.0, 5.4, "Supervisor analogue", [
        "Sees conflicts and refusals.",
        "May ignore a recommendation.",
        "Cannot grant safety from an email.",
    ])
    _card(s, 4.65, 1.25, 4.0, 5.4, "Safety analogue", [
        "Zone release / cert waiver is DENY",
        "without a structured Approval record",
        "(that record is ABSENT → fail closed).",
    ], WARN)
    _card(s, 8.85, 1.25, 4.0, 5.4, "FDE / demo operator", [
        "Runs dashboard + evals.",
        "Shows three frozen stories.",
        "Does not dispatch missions.",
    ], ACCENT)

    # 6 in out
    s = prs.slides.add_slide(blank)
    _slide_bg(s)
    _title(s, "5. In scope / out of scope")
    _table(
        s,
        0.45,
        1.25,
        12.4,
        [
            ["In scope (virtual)", "Out of scope"],
            ["Identity collisions (BOT-COLLISION-*, AMR-044 unmatched)", "Silent merge of two robots"],
            ["Eligibility G1–G8; filter-then-score allocator", "choose_robot as the product recommendation"],
            ["Inventory triple + uncertain", "Single available_qty / wiping sources"],
            ["Cutoff awareness; miss-cutoff recommend", "Release RESTRICTED zone to hit Carrier-A"],
            ["Inject replay in-memory (charger / dock / AS/RS)", "Rewriting CSVs; live plant chaos"],
            ["Observe-only API + dashboard (GET)", "POST fleet, enable physical_control"],
        ],
        [6.4, 6.0],
    )

    # 7 requirements
    s = prs.slides.add_slide(blank)
    _slide_bg(s)
    _title(s, "6. Functional requirements — three workflows")
    _card(s, 0.45, 1.2, 4.0, 5.5, "WF-1 Robot assign", [
        "FR-1 Show eligibility, not HEALTHY.",
        "FR-2 RBT-0020 EXPIRED → INELIGIBLE.",
        "FR-3 RBT-0001 open WO → INELIGIBLE.",
        "FR-4 Assign control is blocked.",
        "FR-5 Alias collision keeps both IDs.",
    ], ACCENT)
    _card(s, 4.65, 1.2, 4.0, 5.5, "WF-2 Inventory", [
        "FR-6 SKU-01146 @ DC-01-Z03-B017",
        "    WMS 205 / ERP 205 / vision 202.",
        "FR-7 uncertain=true; keep all three.",
        "FR-8 Pick-as-known is blocked.",
        "FR-9 Do not copy WMS into vision.",
    ], WARN)
    _card(s, 8.85, 1.2, 4.0, 5.5, "WF-3 Cutoff / complete", [
        "FR-10 ORD-000004 not on_time.",
        "FR-11 Miss cutoff ALLOW (advice).",
        "FR-12 release_zone DENY (G7).",
        "FR-13 ORD-000968 not completable",
        "    while TSK-000968-1 EXECUTING.",
    ], DENY)

    # 8 gates
    s = prs.slides.add_slide(blank)
    _slide_bg(s)
    _title(s, "7. Hard gates (must never flip for a demo)")
    _table(
        s,
        0.45,
        1.2,
        12.4,
        [
            ["Gate", "If true", "Product must"],
            ["G1", "Cert EXPIRED or missing", "INELIGIBLE — do not assign"],
            ["G2", "CMMS OPEN / IN_PROGRESS", "INELIGIBLE — email is not Approval"],
            ["G3", "Payload too small (when known)", "INELIGIBLE"],
            ["G4", "RESTRICTED / UNKNOWN / blocked zone", "Refuse path"],
            ["G5", "Inventory sources disagree", "UNCERTAIN — no single qty"],
            ["G6", "WES ≠ fleet", "Not complete"],
            ["G7", "Safety intent (zone, speed, e-stop, waiver)", "DENY"],
            ["G8", "Unknown ack / missing payload / intermittent", "ABSTAIN"],
        ],
        [1.4, 5.5, 5.5],
    )

    # 9 fixtures
    s = prs.slides.add_slide(blank)
    _slide_bg(s)
    _title(s, "8. Frozen fixtures (do not invent new DCs)")
    _table(
        s,
        0.45,
        1.2,
        12.4,
        [
            ["ID", "Where to look", "Expected UI / API"],
            ["RBT-0020", "data/raw/robots.csv", "EXPIRED + INTERMITTENT → do not assign"],
            ["RBT-0001 / WO-000380", "robots + maintenance.csv", "Open lidar WO; AVAILABLE is not a pass"],
            ["BOT-COLLISION-01", "robot_aliases.csv", "Collision RBT-0001 vs RBT-0002"],
            ["AMR-044", "cascade_001.json", "Unmatched — do not guess RBT-0044"],
            ["SKU-01146", "inventory_snapshot.csv line 2", "DC-01-Z03-B017  205 / 205 / 202"],
            ["ORD-000004", "orders.csv + shipments", "Carrier-A DELAYED; not on-time"],
            ["ORD-000968", "orders + TSK-000968-1", "SHIPPED vs pack EXECUTING"],
        ],
        [3.2, 4.2, 5.0],
    )

    # 10 success
    s = prs.slides.add_slide(blank)
    _slide_bg(s)
    _title(s, "9. Success metrics")
    _subtitle(s, "After = how UC-1 treats the same rows. CSVs are not cleaned.")
    _table(
        s,
        0.45,
        1.25,
        12.4,
        [
            ["Metric", "Before (estate file)", "After (product treatment)"],
            ["False-available treated eligible", "203 / 712 look usable", "0 / 203 treated eligible"],
            ["Invented inventory qty", "legacy_available_qty from WMS", "0 unlabeled available_qty"],
            ["False order complete", "OMS SHIPPED + task EXECUTING", "ORD-000968 completable=false"],
            ["Expired-cert assigns (UC-1)", "Possible via choose_robot", "0 on fixtures / evals"],
            ["EVAL-001–022", "Legacy FAIL where designed", "22 PASS / 0 FAIL on UC-1"],
            ["Physical control", "disabled", "disabled"],
        ],
        [3.6, 4.5, 4.3],
    )

    # 11 architecture
    s = prs.slides.add_slide(blank)
    _slide_bg(s)
    _title(s, "10. Architecture (virtual)")
    _card(s, 0.45, 1.2, 3.0, 5.4, "Observe T0", [
        "CSV / SQLite read",
        "GET /diagnostics",
        "Dashboard KPIs",
        "Conflict lists",
    ], ACCENT)
    _card(s, 3.6, 1.2, 3.0, 5.4, "Refuse T1", [
        "EligibilityPolicy",
        "Inventory UNCERTAIN",
        "TaskReconciler",
        "CutoffAwareness",
    ], WARN)
    _card(s, 6.75, 1.2, 3.0, 5.4, "Preview only", [
        "Filter-then-score",
        "DecisionEngine",
        "ALLOW miss-cutoff",
        "applied=false",
    ], OK)
    _card(s, 9.9, 1.2, 2.95, 5.4, "Never T5", [
        "Fleet POST",
        "WMS write",
        "Motion / PLC",
        "physical_control on",
    ], DENY)

    # 12 UX
    s = prs.slides.add_slide(blank)
    _slide_bg(s)
    _title(s, "11. UX requirements (dashboard / Replit)")
    _bullets(
        s,
        [
            "UX-1 Happy-path buttons Assign / Pick as known / Release zone are blocked or DENY — never look dispatched.",
            "UX-2 Robot screen shows gate reasons (G1 expired cert, G2 open WO), not a green AVAILABLE badge alone.",
            "UX-3 Inventory screen lists WMS / ERP / vision together when they disagree.",
            "UX-4 Banner always shows physical_control: disabled.",
            "UX-5 Notes such as “ignore expired cert” must not change eligibility.",
            "UX-6 One live surface for the client demo (FastAPI http://127.0.0.1:8000/ or Replit using the same GETs — not both competing).",
            "Run: PYTHONPATH=src  python -m uvicorn warehouse_control.api:app --reload",
        ],
        size=15,
    )

    # 13 nfr
    s = prs.slides.add_slide(blank)
    _slide_bg(s)
    _title(s, "12. Non-functional and release")
    _card(s, 0.45, 1.2, 6.1, 5.5, "NFR", [
        "Local Python 3.11+ / 3.12; pytest.",
        "No Azure / cloud OT required.",
        "GET-only product surface; CORS GET ok.",
        "Evals before any new “action-looking” API.",
        "TRACEABILITY keeps FAIL / NOT PROVEN visible.",
        "3 legacy xfails kept (choose_robot baseline).",
    ], ACCENT)
    _card(s, 6.75, 1.2, 6.1, 5.5, "Release (virtual product)", [
        "MAY ship: Repo 2 proof + dashboard.",
        "MUST NOT ship: live OT / customer plant.",
        "Kill: LLM on assign; clean data/; 84% KPI;",
        "“modernization complete.”",
        "Client week: 8 slides + 12 min live + Q&A.",
        "Repo 3 (10 prompts) only if newly mandated.",
    ], WARN)

    # 14 risks
    s = prs.slides.add_slide(blank)
    _slide_bg(s)
    _title(s, "13. Risks and must-nots")
    _table(
        s,
        0.45,
        1.2,
        12.4,
        [
            ["Risk", "If we fail", "Control"],
            ["Unsafe assign", "Expired / open-CMMS robot recommended", "G1 G2; UI block Assign"],
            ["Invented qty", "Wrong pick in the story", "Keep 205/205/202; G5"],
            ["Safety for on-time", "Zone release under Carrier-A pressure", "G7 DENY"],
            ["Injection", "Chat/email becomes Policy", "Ignore shadow_note; EVAL-020"],
            ["Fake ROI", "Invented dollars", "No cost field — say so"],
            ["Demo looks live", "Client thinks robots will move", "Disabled banner; no POST"],
        ],
        [2.4, 5.0, 5.0],
    )

    # 15 ask
    s = prs.slides.add_slide(blank)
    _slide_bg(s)
    _title(s, "14. Ask")
    _card(
        s,
        0.45,
        1.3,
        12.4,
        5.4,
        "We ask the client to accept",
        [
            "1.  This virtual control tower as the decision layer proof (observe / refuse / explain).",
            "2.  Deterministic gates for safety and inventory — not an agent dispatcher.",
            "3.  Physical control remains disabled; no live warehouse in this engagement.",
            "4.  Value = avoided false work and avoided unsafe assign. Financial ROI is a later measurement plan, not a number we will invent.",
            "5.  Optional next: production-shaped packaging (same engine, one UI, CI) still fully virtual.",
        ],
        OK,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(path)
