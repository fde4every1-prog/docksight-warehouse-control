"""DockSight final stakeholder deck — light warehouse theme."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

OUT = (
    Path(__file__).resolve().parents[1]
    / "docs"
    / "DockSight_Final_Presentation.pptx"
)

# Light warehouse: kraft paper, steel rack, aisle-tape amber
PAPER = RGBColor(0xF6, 0xF3, 0xEC)
PAPER_DEEP = RGBColor(0xEB, 0xE6, 0xDA)
CARD = RGBColor(0xFF, 0xFC, 0xF7)
INK = RGBColor(0x2B, 0x33, 0x38)
MUTED = RGBColor(0x5C, 0x65, 0x6C)
STEEL = RGBColor(0x3D, 0x5A, 0x6C)
STEEL_DEEP = RGBColor(0x2F, 0x45, 0x54)
AMBER = RGBColor(0xC9, 0x96, 0x2A)
AMBER_SOFT = RGBColor(0xF4, 0xE8, 0xC6)
LINE = RGBColor(0xD6, 0xCE, 0xBE)
STOP = RGBColor(0xA8, 0x4B, 0x40)
GO = RGBColor(0x3E, 0x6F, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _set_run(run, text, size=16, bold=False, color=INK):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _send_back(slide, shape):
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def _notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def _slide_chrome(slide, footer):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    _fill(bg, PAPER)
    _send_back(slide, bg)

    rack = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.10), Inches(7.5))
    _fill(rack, STEEL)

    tape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Inches(7.08), Inches(13.333), Inches(0.10)
    )
    _fill(tape, AMBER)

    foot = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, Inches(7.18), Inches(13.333), Inches(0.32)
    )
    _fill(foot, STEEL_DEEP)

    tb = slide.shapes.add_textbox(Inches(0.42), Inches(7.20), Inches(12.5), Inches(0.26))
    r = tb.text_frame.paragraphs[0].add_run()
    _set_run(r, footer, 11, False, PAPER)

    # Quiet crate mark — three stacked cartons, not clipart
    for i, (w, h) in enumerate(((0.38, 0.22), (0.32, 0.18), (0.26, 0.14))):
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(12.55 + i * 0.04),
            Inches(0.18 + i * 0.08),
            Inches(w),
            Inches(h),
        )
        box.fill.solid()
        box.fill.fore_color.rgb = AMBER_SOFT if i else PAPER_DEEP
        box.line.color.rgb = LINE


def _eyebrow(slide, text):
    box = slide.shapes.add_textbox(Inches(0.42), Inches(0.16), Inches(11.8), Inches(0.28))
    r = box.text_frame.paragraphs[0].add_run()
    _set_run(r, text, 11, True, STEEL)


def _title(slide, text, y=0.42):
    box = slide.shapes.add_textbox(Inches(0.42), Inches(y), Inches(12.4), Inches(0.48))
    tf = box.text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    _set_run(r, text, 26, True, INK)


def _subtitle(slide, text, y=0.88):
    box = slide.shapes.add_textbox(Inches(0.42), Inches(y), Inches(12.4), Inches(0.42))
    tf = box.text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    _set_run(r, text, 15, False, MUTED)


def _card(slide, left, top, w, h, heading, lines, heading_color=STEEL):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(h)
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = CARD
    shp.line.color.rgb = LINE
    try:
        shp.adjustments[0] = 0.08
    except Exception:
        pass

    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(0.08), Inches(h)
    )
    _fill(bar, heading_color)

    tb = slide.shapes.add_textbox(
        Inches(left + 0.22), Inches(top + 0.14), Inches(w - 0.38), Inches(h - 0.26)
    )
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    _set_run(r, heading, 14, True, heading_color)
    for line in lines:
        p = tf.add_paragraph()
        p.space_before = Pt(7)
        r = p.add_run()
        _set_run(r, line, 13, False, INK)


def _table(slide, left, top, width, rows, col_w, row_h=0.42):
    table_shape = slide.shapes.add_table(
        len(rows), len(rows[0]), Inches(left), Inches(top), Inches(width), Inches(row_h * len(rows))
    )
    table = table_shape.table
    for i, w in enumerate(col_w):
        table.columns[i].width = Inches(w)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            cell.text = ""
            cell.fill.solid()
            if ri == 0:
                cell.fill.fore_color.rgb = STEEL
                color, bold, size = WHITE, True, 12
            else:
                cell.fill.fore_color.rgb = CARD if ri % 2 else PAPER
                color, bold, size = INK, False, 12
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            _set_run(r, val, size, bold, color)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    foot_std = "DockSight  ·  Virtual warehouse control tower  ·  Synthetic proof  ·  Not live robots"

    # ------------------------------------------------------------------ 1
    s = prs.slides.add_slide(blank)
    _slide_chrome(s, foot_std)
    _eyebrow(s, "DOCKSIGHT  ·  FDE CAPSTONE  ·  SLIDE 1")
    _title(s, "The business case")
    _subtitle(
        s,
        "The estate already paid for robots and warehouse systems. Exceptions still create late trucks, false completes, and unsafe assigns.",
    )
    _card(
        s,
        0.42,
        1.38,
        4.05,
        5.52,
        "Problem statement",
        [
            "This is not a greenfield warehouse.",
            "OMS, WMS, WES, fleet, CMMS, vision, and TMS already run — and they often do not agree.",
            "Supervisors still hit Go because no one screen says stop.",
            "The business case is not “buy AI.”",
            "It is: stop a bad go-ahead when the computer picture is not the floor.",
        ],
        STEEL,
    )
    _card(
        s,
        4.62,
        1.38,
        4.05,
        5.52,
        "Exact real-world problem",
        [
            "A supervisor must send a robot to pick an order before the truck cutoff.",
            "A robot looks free — cert expired, or a repair is still open.",
            "The same bin shows three quantities.",
            "The order looks shipped while the job is still running — or TMS is already DELAYED.",
            "Two machines can share one radio nickname.",
            "People paper over this with emails and unofficial lists.",
        ],
        STOP,
    )
    _card(
        s,
        8.82,
        1.38,
        4.08,
        5.52,
        "Solution — DockSight",
        [
            "A virtual watch-and-stop control tower.",
            "One path: order → stock → job → robot → zone → truck.",
            "ALLOW only when sources agree and the robot is eligible.",
            "Else DENY or wait. Assign, pick, and zone release stay blocked.",
            "Rules first. Copilot later. No LLM on the write path.",
            "Observe only. Physical control stays off.",
        ],
        GO,
    )
    _notes(
        s,
        "Say: The warehouse already has robots. The systems do not agree. People still hit Go. "
        "We are not making robots faster. We are stopping a bad assign when digital state is not physical truth. "
        "DockSight is a watch-and-stop tower — not live control.",
    )

    # ------------------------------------------------------------------ 2
    s = prs.slides.add_slide(blank)
    _slide_chrome(s, foot_std)
    _eyebrow(s, "EVIDENCE  ·  BROWNFIELD SNAPSHOT  ·  DO NOT CLEAN THE FILES")
    _title(s, "Why this is a business problem, not a demo bug")
    _subtitle(s, "Inherited estate. Competing systems of record. Exceptions are the mass, not the tail.")
    _table(
        s,
        0.42,
        1.38,
        12.48,
        [
            ["What we measured", "Snapshot", "What happens on the floor"],
            [
                "Robot looks free, should not move",
                "203 / 712  (28.5%)",
                "Expired cert or open repair — still pickable on the old path",
            ],
            [
                "Same SKU, three quantities",
                "7,382 / 9,360  (78.9%)",
                "Pick that is not really there",
            ],
            [
                "Job status split (WES ≠ fleet)",
                "7,351 / 8,732  (84.2%)",
                "Order can look done while the robot is still working",
            ],
            [
                "Order status split (OMS ≠ WMS)",
                "1,283 / 3,500  (36.7%)",
                "Customer or planner sees a different story than the aisle",
            ],
            [
                "Truck already late (TMS DELAYED)",
                "586 / 3,500  (16.7%)",
                "OMS SLA can still look fine",
            ],
        ],
        [4.3, 2.7, 5.48],
        row_h=0.78,
    )
    _notes(
        s,
        "Do not present teammate 62% on-time or 0.21% inventory as this estate. "
        "Do not use 84% as-of cutoff as a board KPI. "
        "Say: we did not clean the books. The messy rows stay messy on purpose.",
    )

    # ------------------------------------------------------------------ 3
    s = prs.slides.add_slide(blank)
    _slide_chrome(s, foot_std)
    _eyebrow(s, "USER  ·  DECISION  ·  SUCCESS")
    _title(s, "Who we help, and what “done” looks like")
    _subtitle(s, "Primary user is the floor supervisor analogue — not an autonomous dispatcher.")
    _card(
        s,
        0.42,
        1.38,
        4.05,
        5.52,
        "Primary user",
        [
            "Warehouse supervisor / exception handler.",
            "They already open five systems and still guess.",
            "Safety officer is a second voice: zone and cert are not an email.",
            "We do not replace the operator. We give them a stop they can defend.",
        ],
        STEEL,
    )
    _card(
        s,
        4.62,
        1.38,
        4.05,
        5.52,
        "The decision",
        [
            "May we assign this robot to this order, for this stock, toward this truck?",
            "Or must we wait?",
            "Legal answers: ALLOW · DENY · wait (abstain).",
            "A chatbot that always recommends a robot is the wrong product.",
        ],
        AMBER,
    )
    _card(
        s,
        8.82,
        1.38,
        4.08,
        5.52,
        "Success at the end of the path",
        [
            "They see why it is stop or go — without five screens.",
            "Wrong robot is not sent.",
            "Uncertain stock is not picked as known.",
            "Customer is not told the truck left when TMS is DELAYED.",
            "Unsafe assign in the fixtures we show = 0.",
        ],
        GO,
    )
    _notes(
        s,
        "If asked who the AI user is: there is no AI user on the write path. The supervisor is the user. Copilot is optional explain-only and default off.",
    )

    # ------------------------------------------------------------------ 4
    s = prs.slides.add_slide(blank)
    _slide_chrome(s, foot_std)
    _eyebrow(s, "TODAY  ·  WHERE IT BREAKS")
    _title(s, "Current process — and where humans enter")
    _subtitle(s, "The intended chain is OMS → WMS → WES → fleet → pack → TMS. In this snapshot, zero orders close that chain cleanly.")
    _card(
        s,
        0.42,
        1.38,
        6.15,
        5.52,
        "Without DockSight",
        [
            "Look up the order in OMS. Trust WMS quantity. Ask fleet who is AVAILABLE.",
            "Highest battery or a local score often wins.",
            "If something looks wrong: email, radio, unofficial wave list (FINAL_v7).",
            "Human enters late — after a bad assign, a missed cutoff, or a safety event.",
            "Override is not a recorded approval. It is a message that says “I approved it.”",
        ],
        STOP,
    )
    _card(
        s,
        6.77,
        1.38,
        6.12,
        5.52,
        "With DockSight",
        [
            "Human enters at the gate — before Go.",
            "Identity, eligibility, inventory disagreement, job split, zone, cutoff are visible first.",
            "If evidence is missing or sources conflict: wait. Do not invent true_qty.",
            "Assign / Pick / Release controls stay blocked on this proof.",
            "Shadow emails stay evidence of the old workaround — not policy.",
        ],
        GO,
    )
    _notes(
        s,
        "Closest inherited “success” ORD-000968 is already departed while pack-feed is still EXECUTING. That is false complete — use it if they ask for an example.",
    )

    # ------------------------------------------------------------------ 5
    s = prs.slides.add_slide(blank)
    _slide_chrome(s, foot_std)
    _eyebrow(s, "WHAT WE BUILT  ·  OBSERVE ONLY")
    _title(s, "The path we walk in the demo")
    _subtitle(s, "Same six stages for a clean order and a broken order. We did not rewrite the warehouse.")

    stages = [
        ("1  Order", "OMS vs WMS\nDo they agree?"),
        ("2  Stock", "WMS · ERP · vision\nThree claims, not one qty"),
        ("3  Job", "WES vs fleet\nComplete only if both"),
        ("4  Robot", "Eligible, not AVAILABLE\nCert · repair · site"),
        ("5  Zone", "RESTRICTED stays closed\nNo demo waiver"),
        ("6  Truck", "TMS + actual time\nDELAYED means empty actual"),
    ]
    x = 0.42
    for title, body in stages:
        stage = s.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.42), Inches(1.95), Inches(3.35)
        )
        stage.fill.solid()
        stage.fill.fore_color.rgb = CARD
        stage.line.color.rgb = LINE
        tb = s.shapes.add_textbox(Inches(x + 0.10), Inches(1.55), Inches(1.75), Inches(3.05))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        _set_run(r, title, 14, True, STEEL)
        for line in body.split("\n"):
            p = tf.add_paragraph()
            p.space_before = Pt(10)
            r = p.add_run()
            _set_run(r, line, 12, False, INK)
        x += 2.12

    _card(
        s,
        0.42,
        4.95,
        12.48,
        1.95,
        "What this product is",
        [
            "A join of names that already exist: warehouse, robot, order, task, sku+bin, shipment. That is the ontology — meaning and keys, not a graph database.",
            "When they disagree we keep both sides as a Conflict. We do not pick a winner system because its name sounds official.",
            "GET-only API. physical_control stays disabled. We are not driving robots from this screen.",
        ],
        STEEL,
    )
    _notes(
        s,
        "If asked about ontology: it is the dictionary so we can join without inventing truth. Not OWL, not Neo4j on the decision path.",
    )

    # ------------------------------------------------------------------ 6
    s = prs.slides.add_slide(blank)
    _slide_chrome(s, foot_std)
    _eyebrow(s, "DEMO  ·  TWO STORIES  ·  DO NOT INVENT NEW IDS")
    _title(s, "What we will show live")
    _subtitle(s, "One agreed path. One inherited failure. Same screen. Hide raw JSON unless they ask.")
    _card(
        s,
        0.42,
        1.38,
        6.15,
        5.52,
        "Green — ORD-009999  ·  RBT-0012",
        [
            "OMS and WMS agree. Qty 40 / 40 / 40.",
            "One task, WES = fleet COMPLETE.",
            "Robot eligible. Zone not restricted.",
            "TMS DEPARTED with an actual time.",
            "This is the control: “agreed” is possible.",
            "We added this slice. We did not clean the rest of the books.",
        ],
        GO,
    )
    _card(
        s,
        6.77,
        1.38,
        6.12,
        5.52,
        "Broken — ORD-000004  ·  watch RBT-0020 / RBT-0644",
        [
            "OMS STAGED, WMS EXCEPTION, TMS DELAYED, actual blank.",
            "SKU-01146 in DC-01: quantities do not match.",
            "TSK-000004-1: WES ≠ fleet. Robot can be unsafe or wrong DC.",
            "Assign stays blocked. Pick-as-known stays blocked.",
            "Also in the pack if asked: ORD-000968 false complete; BOT-COLLISION alias.",
        ],
        STOP,
    )
    _notes(
        s,
        "Open /workflow. Select ORD-009999 then ORD-000004. Do not click Assign. If they ask contention: same location / dock / sku is a conflict, not a physical store. "
        "If they ask teammate 7000 rows: those are metrics volume, not this failure story.",
    )

    # ------------------------------------------------------------------ 7
    s = prs.slides.add_slide(blank)
    _slide_chrome(s, foot_std)
    _eyebrow(s, "WHAT FDE CHANGED")
    _title(s, "Before this training  ·  after this capstone")
    _subtitle(s, "The skill we are presenting: evidence before orchestration.")
    _card(
        s,
        0.42,
        1.38,
        6.15,
        5.52,
        "Before FDE — what I might have done",
        [
            "Clean the CSVs and pick one source of truth.",
            "Dashboard of on-time % and utilization.",
            "Put a chatbot on WMS and call it modernization.",
            "Trust AVAILABLE and send the robot.",
            "Build a graph or twin first.",
            "Invent a fully green dataset so the demo never fails.",
        ],
        MUTED,
    )
    _card(
        s,
        6.77,
        1.38,
        6.12,
        5.52,
        "After FDE — what we did instead",
        [
            "Named journeys, not a cleaned average.",
            "Digital ≠ physical. Available ≠ suitable ≠ allowed.",
            "Keep both sides of a conflict. No invented true_qty.",
            "Stop is a valid answer. LLM off the write path.",
            "Ontology = words and join keys, not a product we bought.",
            "Evals before anything that looks autonomous.",
        ],
        GO,
    )
    _notes(
        s,
        "Close this slide with: before = hide the mess and add AI. After = show the mess and stop the unsafe Go.",
    )

    # ------------------------------------------------------------------ 8
    s = prs.slides.add_slide(blank)
    _slide_chrome(s, foot_std)
    _eyebrow(s, "BOUNDARIES  ·  SAY THIS OUT LOUD")
    _title(s, "What this proof is — and is not")
    _card(
        s,
        0.42,
        1.38,
        6.15,
        5.52,
        "In this room we can claim",
        [
            "The problem is evidenced in the inherited files.",
            "A safe control idea: watch, refuse, recommend miss-cutoff.",
            "Three workflows: assign, inventory uncertainty, cutoff.",
            "Identity collisions stay visible. AMR-044 stays unmatched.",
            "The warehouse is not modernized yet. That is honest.",
        ],
        GO,
    )
    _card(
        s,
        6.77,
        1.38,
        6.12,
        5.52,
        "We will not claim",
        [
            "Live robots, PLC, or WMS write-back.",
            "Agents that recover the plant without a human.",
            "Dollar ROI or headcount reduction (no cost field).",
            "ISO 42001. Azure. A mandatory knowledge graph.",
            "Teammate cycle-time as this estate’s on-time KPI.",
            "Turning physical_control on for a nicer demo.",
        ],
        STOP,
    )
    _notes(
        s,
        "If they ask MTTR: not computable from this snapshot. Auto-refuse is not plant recovery. "
        "If they ask global vs local: eligibility is global policy; legacy_score is local greed. We keep lexicographic safety > truth > SLA.",
    )

    # ------------------------------------------------------------------ 9
    s = prs.slides.add_slide(blank)
    _slide_chrome(s, "Clock starts only if you accept this proof  ·  Robots stay off until a separate safety case")
    _eyebrow(s, "ASK  ·  NEXT 90 DAYS")
    _title(s, "What we need from you")
    _subtitle(s, "This proof shows the problem and a safe control idea. Production-ready is a later mandate.")
    _card(
        s,
        0.42,
        1.38,
        4.05,
        5.52,
        "Days 1–30",
        [
            "Same proof, real operator path.",
            "Stop showing “highest battery” as the picker.",
            "Always show who the robot really is.",
            "Honest KPIs — no fake 84% on-time.",
            "Still synthetic. Still no live fleet.",
        ],
        STEEL,
    )
    _card(
        s,
        4.62,
        1.38,
        4.05,
        5.52,
        "Days 31–60",
        [
            "Optional helper that explains a stop.",
            "It cannot send a robot.",
            "When qty disagrees: people count the bin.",
            "Agree clocks and mission IDs for a future feed.",
        ],
        AMBER,
    )
    _card(
        s,
        8.82,
        1.38,
        4.08,
        5.52,
        "Days 61–90",
        [
            "Train exception staff on this tower.",
            "Measure on a live slice only with new consent.",
            "Only then discuss whether control could ever turn on — with a safety case.",
        ],
        GO,
    )
    _notes(
        s,
        "Ask: accept this as the problem and the control idea. Do not turn on robots for the showcase. Start the 90-day track only if they say so.",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(path)
