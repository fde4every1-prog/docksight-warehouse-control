"""Build P01 current-architecture PPTX. Does not change baseline 2.0.0."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "artifacts" / "P01-v1-CURRENT_ARCHITECTURE.pptx"

NAVY = RGBColor(0x0F, 0x2C, 0x4C)
NAVY_MID = RGBColor(0x1B, 0x4F, 0x72)
TEAL = RGBColor(0x1A, 0x6B, 0x6B)
AMBER = RGBColor(0xB8, 0x6E, 0x00)
RED = RGBColor(0x8B, 0x2E, 0x2E)
SLATE = RGBColor(0x5C, 0x67, 0x73)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFFWHITE = RGBColor(0xF4, 0xF6, 0xF8)
INK = RGBColor(0x1A, 0x1D, 0x21)


def set_run(run, text, size=14, bold=False, color=INK, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def box(slide, l, t, w, h, fill, text, size=11, bold=True, font_color=WHITE, align=PP_ALIGN.CENTER):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    fill_shape(sh, fill)
    sh.adjustments[0] = 0.08
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.06)
    tf.margin_bottom = Inches(0.06)
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(2)
        run = p.add_run()
        set_run(run, line, size, bold, font_color)
    return sh


def add_textbox(slide, l, t, w, h, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    set_run(run, text, size, bold, color)
    return tb


def add_multiline(slide, l, t, w, h, lines, size=14, color=INK, bold_first=False):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(6)
        run = p.add_run()
        set_run(run, line, size, bold=(bold_first and i == 0), color=color)
    return tb


def footer(slide, page, total=10):
    add_textbox(
        slide,
        Inches(0.5),
        Inches(7.05),
        Inches(10.5),
        Inches(0.3),
        f"P01-v1  |  Baseline 2.0.0 frozen  |  WRCC current architecture  |  {page}/{total}",
        size=10,
        color=SLATE,
    )


def header_bar(slide, title):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.92))
    fill_shape(bar, NAVY)
    add_textbox(slide, Inches(0.5), Inches(0.22), Inches(12), Inches(0.5), title, size=24, bold=True, color=WHITE)


def blank_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    fill_shape(bg, OFFWHITE)
    spTree = slide.shapes._spTree
    sp = bg._element
    spTree.remove(sp)
    spTree.insert(2, sp)
    return slide


def arrow_row(slide, labels, top, fill=NAVY_MID, width=1.45, height=0.58, start_left=0.4):
    left = start_left
    gap = 0.18
    for i, label in enumerate(labels):
        box(slide, Inches(left), Inches(top), Inches(width), Inches(height), fill, label, size=10)
        left += width
        if i < len(labels) - 1:
            add_textbox(
                slide,
                Inches(left),
                Inches(top + 0.08),
                Inches(gap),
                Inches(0.4),
                "→",
                size=16,
                bold=True,
                color=NAVY,
                align=PP_ALIGN.CENTER,
            )
            left += gap
    return left


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 title
    s = blank_slide(prs)
    hero = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    fill_shape(hero, NAVY)
    add_textbox(s, Inches(0.7), Inches(2.0), Inches(12), Inches(0.5), "PROMPT 01  ·  FORENSICS", size=14, bold=True, color=RGBColor(0xC5, 0xD8, 0xE8))
    add_textbox(s, Inches(0.7), Inches(2.5), Inches(12), Inches(1.2), "Current-state architecture", size=40, bold=True, color=WHITE)
    add_textbox(
        s,
        Inches(0.7),
        Inches(3.7),
        Inches(11.5),
        Inches(1.0),
        "AI-Native Autonomous Warehouse & Robotics Control Center\nExpanded from docs/03  ·  Baseline v2.0.0 frozen  ·  Overlay 0.1.0",
        size=16,
        color=RGBColor(0xD6, 0xE2, 0xEA),
    )
    add_textbox(s, Inches(0.7), Inches(6.6), Inches(11), Inches(0.4), "No component has complete warehouse truth.", size=16, bold=True, color=AMBER)

    # 2 tensions
    s = blank_slide(prs)
    header_bar(s, "What this architecture is proving")
    box(s, Inches(0.5), Inches(1.25), Inches(6.0), Inches(2.35), NAVY_MID, "Digital warehouse state\n≠\nPhysical warehouse state", size=18)
    box(s, Inches(6.8), Inches(1.25), Inches(6.0), Inches(2.35), TEAL, "Robot available\n≠ suitable ≠ optimal", size=18)
    add_multiline(
        s,
        Inches(0.55),
        Inches(3.85),
        Inches(12.2),
        Inches(2.8),
        [
            "This is a brownfield estate of 18 DCs. ERP, OMS, WMS, WES, fleet/WCS, TMS, CMMS, vision, labor and emails evolved independently.",
            "The happy path usually works. Exceptions are reconciled by humans (shadow email/Excel), not by a single system of record.",
            "This repo is a synthetic snapshot + observe-only API. It is not a live digital twin and does not command robots.",
        ],
        size=16,
    )
    footer(s, 2)

    # 3 original sketch
    s = blank_slide(prs)
    header_bar(s, "Inherited sketch  ·  docs/03_current_state_architecture.md")
    add_textbox(s, Inches(0.5), Inches(1.15), Inches(12), Inches(0.35), "Frozen baseline — not edited. This deck expands it with repo evidence.", size=13, color=SLATE)
    sketch = [
        "ERP → OMS → WMS → WES → Fleet Managers / WCS → Robots / PLCs / Conveyors / ASRS",
        "                              ↘ Vision / Safety / IoT",
        "CMMS ↔ Fleet                       ↓",
        "Labor Mgmt                    Physical warehouse",
        "TMS ← staging / ship confirmation",
        "",
        "Shadow layer: CSV/Excel-like exports + emails + radio / manual overrides",
    ]
    panel = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.6), Inches(12.3), Inches(4.6))
    fill_shape(panel, WHITE)
    panel.line.color.rgb = RGBColor(0xD0, 0xD7, 0xDE)
    add_multiline(s, Inches(0.75), Inches(1.85), Inches(11.8), Inches(4.2), sketch, size=16, color=NAVY)
    footer(s, 3)

    # 4 documented flow
    s = blank_slide(prs)
    header_bar(s, "Documented as-is flow  (same boxes as docs/03)")
    add_textbox(s, Inches(0.5), Inches(1.1), Inches(12), Inches(0.3), "Happy-path order movement. Dashed conceptually: shadow ops and CMMS↔fleet.", size=13, color=SLATE)
    arrow_row(s, ["ERP", "OMS", "WMS", "WES", "Fleet", "Robots", "Physical", "TMS"], top=1.6, fill=NAVY_MID, width=1.38, start_left=0.45)
    add_textbox(s, Inches(4.4), Inches(2.35), Inches(4.5), Inches(0.3), "WES also drives machines and sensing", size=12, color=SLATE, align=PP_ALIGN.CENTER)
    box(s, Inches(3.55), Inches(2.7), Inches(1.7), Inches(0.55), TEAL, "WCS / PLC", size=11)
    box(s, Inches(5.45), Inches(2.7), Inches(2.35), Inches(0.55), TEAL, "Conveyor / ASRS / dock", size=10)
    box(s, Inches(8.05), Inches(2.7), Inches(1.55), Inches(0.55), TEAL, "Vision", size=11)
    box(s, Inches(9.8), Inches(2.7), Inches(1.55), Inches(0.55), TEAL, "Safety / IoT", size=11)
    add_textbox(s, Inches(0.5), Inches(3.45), Inches(12), Inches(0.3), "People, maintenance, unofficial process", size=13, bold=True, color=NAVY)
    box(s, Inches(0.5), Inches(3.85), Inches(2.4), Inches(0.7), AMBER, "CMMS  ↔  Fleet", size=12, font_color=WHITE)
    box(s, Inches(3.15), Inches(3.85), Inches(2.2), Inches(0.7), AMBER, "Labor mgmt", size=12, font_color=WHITE)
    box(s, Inches(5.6), Inches(3.85), Inches(7.2), Inches(0.7), RED, "Shadow: emails / Excel / radio  →  OMS, WMS, Fleet", size=12)
    add_multiline(
        s,
        Inches(0.5),
        Inches(4.8),
        Inches(12.3),
        Inches(1.9),
        [
            "Enterprise: ERP qty, OMS order + carrier_cutoff",
            "Warehouse software: WMS bins/qty/status, WES task.wes_status",
            "Floor: fleet, WCS/PLC, robots, conveyors, ASRS, chargers, docks → physical warehouse → TMS",
        ],
        size=14,
    )
    footer(s, 4)

    # 5 repo implementation
    s = blank_slide(prs)
    header_bar(s, "What this repo actually runs  ·  thin control plane")
    add_textbox(s, Inches(0.5), Inches(1.1), Inches(12.3), Inches(0.35), "No live WarehouseState. No DecisionEngine. No robot write path. State = snapshots.", size=14, bold=True, color=RED)
    cols = [
        (NAVY_MID, "data/ snapshots", "orders, inventory, tasks,\nrobots, aliases, zones,\ncharging, telemetry,\nevents, shipments, safety,\nlabor, shadow email/Excel"),
        (TEAL, "Baseline 2.0.0", "diagnostics.py counters\nlegacy allocator\n(battery + ONLINE)\nlegacy inventory\n(trusts WMS)\nSQLite + GET API"),
        (AMBER, "Contracts only", "fleet_api_v1.yaml\nno idempotency key\nfleet_api_v2.yaml\nvehicle_id\nnot wired to FastAPI"),
        (RED, "Absent", "DecisionEngine\nCopilot / RAG\nKeep-out polygons\nTote model\nDigital twin runtime\nspecs/ WRCC contract"),
    ]
    for i, (color, title, body) in enumerate(cols):
        x = 0.4 + i * 3.2
        box(s, Inches(x), Inches(1.6), Inches(3.0), Inches(0.55), color, title, size=13)
        body_box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.25), Inches(3.0), Inches(3.55))
        fill_shape(body_box, WHITE)
        body_box.line.color.rgb = color
        tf = body_box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.12)
        tf.margin_top = Inches(0.12)
        p = tf.paragraphs[0]
        run = p.add_run()
        set_run(run, body, 13, False, INK)
    footer(s, 5)

    # 6 competing records
    s = blank_slide(prs)
    header_bar(s, "Competing systems of record")
    fights = [
        ("Robot identity", "WMS alias vs fleet_id vs CMMS alias", "6 alias collisions"),
        ("Inventory qty", "WMS vs ERP vs vision vs operator", "7,382 disagreements"),
        ("Task / order", "OMS vs WMS vs WES vs fleet", "7,351 WES ≠ fleet"),
        ("Availability", "Fleet AVAILABLE vs CMMS vs cert", "176 + 35 expired-cert online"),
        ("Cutoff / time", "OMS SLA vs email vs event clocks", "Cutoff −35 min; 7,448 skews"),
    ]
    for i, (title, who, ev) in enumerate(fights):
        y = 1.2 + i * 1.05
        box(s, Inches(0.45), Inches(y), Inches(2.6), Inches(0.9), NAVY_MID, title, size=13)
        who_b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.2), Inches(y), Inches(5.7), Inches(0.9))
        fill_shape(who_b, WHITE)
        who_b.line.color.rgb = RGBColor(0xD0, 0xD7, 0xDE)
        tf = who_b.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        set_run(run, who, 14, False, INK)
        box(s, Inches(9.1), Inches(y), Inches(3.7), Inches(0.9), AMBER if i < 4 else TEAL, ev, size=12)
    footer(s, 6)

    # 7 evidence strip
    s = blank_slide(prs)
    header_bar(s, "Baseline diagnostic evidence  ·  VERIFICATION.md")
    stats = [
        ("712", "Robots"),
        ("6", "Alias collisions"),
        ("7,382", "Inventory fights"),
        ("7,351", "WES ≠ fleet"),
        ("176", "CMMS vs AVAILABLE"),
        ("35", "Expired cert, connected"),
    ]
    for i, (n, lab) in enumerate(stats):
        x = 0.45 + (i % 3) * 4.2
        y = 1.4 + (i // 3) * 2.4
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(3.95), Inches(2.15))
        fill_shape(card, WHITE)
        card.line.color.rgb = RGBColor(0xD0, 0xD7, 0xDE)
        add_textbox(s, Inches(x + 0.15), Inches(y + 0.35), Inches(3.65), Inches(0.9), n, size=36, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        add_textbox(s, Inches(x + 0.15), Inches(y + 1.3), Inches(3.65), Inches(0.5), lab, size=14, color=SLATE, align=PP_ALIGN.CENTER)
    footer(s, 7)

    # 8 authority
    s = blank_slide(prs)
    header_bar(s, "Authority boundary today")
    steps = [
        (NAVY_MID, "1. Observe", "diagnostics.py\nGET /health\nGET /diagnostics\nGET /robots/{id}"),
        (TEAL, "2. Recommend", "Not implemented\nin baseline 2.0.0"),
        (AMBER, "3. Human", "Email / supervisor\nsafety docs\nCMMS + labor"),
        (RED, "4. Execute", "Blocked\nphysical_control:\ndisabled"),
    ]
    for i, (color, title, body) in enumerate(steps):
        x = 0.5 + i * 3.2
        box(s, Inches(x), Inches(1.4), Inches(2.95), Inches(0.6), color, title, size=14)
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.15), Inches(2.95), Inches(2.5))
        fill_shape(b, WHITE)
        b.line.color.rgb = color
        tf = b.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.12)
        tf.margin_top = Inches(0.2)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        set_run(run, body, 14, False, INK)
        if i < 3:
            add_textbox(s, Inches(x + 2.85), Inches(2.9), Inches(0.4), Inches(0.4), "→", size=20, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_multiline(
        s,
        Inches(0.5),
        Inches(4.9),
        Inches(12.3),
        Inches(1.7),
        [
            "docs/06: e-stop, safety PLC, and speed-limit changes require human authority.",
            "There is no copilot on the write path — and adding a chatbot on /diagnostics is not modernization.",
            "Runnable proof: python -m warehouse_control.cli diagnostics   ·   uvicorn warehouse_control.api:app",
        ],
        size=15,
    )
    footer(s, 8)

    # 9 named examples
    s = blank_slide(prs)
    header_bar(s, "Named examples you can open in the repo")
    rows = [
        ("Identity", "BOT-COLLISION-01", "CMMS = RBT-0001  ·  WMS = RBT-0002"),
        ("Safety", "RBT-0007", "EXPIRED cert, still ONLINE — legacy allocator would still score it"),
        ("Order", "ORD-000003", "OMS PACKING vs WMS EXCEPTION; all 3 tasks WES ≠ fleet"),
        ("Inventory", "SKU-01146 @ DC-01-Z03-B017", "WMS 205 · ERP 205 · vision 202"),
        ("Shadow", "ops_emails.txt", "Dock 7→Z09 unmapped; camera 18; cutoff −35 min; AMR-044"),
    ]
    box(s, Inches(0.45), Inches(1.2), Inches(2.2), Inches(0.45), NAVY, "Lens", size=11)
    box(s, Inches(2.75), Inches(1.2), Inches(3.5), Inches(0.45), NAVY, "Open this ID", size=11)
    box(s, Inches(6.35), Inches(1.2), Inches(6.5), Inches(0.45), NAVY, "What it shows", size=11)
    for i, (a, b, c) in enumerate(rows):
        y = 1.75 + i * 0.85
        fill = WHITE
        ba = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45), Inches(y), Inches(2.2), Inches(0.75))
        fill_shape(ba, TEAL if i % 2 == 0 else NAVY_MID)
        tf = ba.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        set_run(run, a, 12, True, WHITE)
        for left, w, txt, col in (
            (2.75, 3.5, b, INK),
            (6.35, 6.5, c, INK),
        ):
            sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(y), Inches(w), Inches(0.75))
            fill_shape(sh, fill)
            sh.line.color.rgb = RGBColor(0xD0, 0xD7, 0xDE)
            tf = sh.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            set_run(run, txt, 12, False, col)
    footer(s, 9)

    # 10 close
    s = blank_slide(prs)
    header_bar(s, "How to use this with the rest of P01")
    items = [
        ("docs/03_current_state_architecture.md", "Frozen original sketch"),
        ("P01-v1-CURRENT_ARCHITECTURE.md", "Mermaid source for this deck"),
        ("P01-v1-CURRENT_STATE.md", "File-level inventory"),
        ("P01-v1-SDD_GAP_ANALYSIS.md", "WRCC PASS / PARTIAL / FAIL / NOT PROVEN"),
        ("P01-v1-MODERNIZATION_BACKLOG.md", "What we will overlay — not rewrite"),
    ]
    for i, (name, role) in enumerate(items):
        y = 1.25 + i * 0.85
        box(s, Inches(0.5), Inches(y), Inches(7.4), Inches(0.7), NAVY_MID, name, size=14)
        box(s, Inches(8.1), Inches(y), Inches(4.7), Inches(0.7), TEAL, role, size=13)
    footer(s, 10)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
