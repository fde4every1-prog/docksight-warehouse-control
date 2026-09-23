"""Simple 3-slide current-architecture deck."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parents[1] / "artifacts" / "P01-v1-CURRENT_ARCHITECTURE-SIMPLE.pptx"

NAVY = RGBColor(0x0F, 0x2C, 0x4C)
BLUE = RGBColor(0x1B, 0x4F, 0x72)
TEAL = RGBColor(0x1A, 0x6B, 0x6B)
AMBER = RGBColor(0xB8, 0x6E, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1A, 0x1D, 0x21)
SLATE = RGBColor(0x5C, 0x67, 0x73)
BG = RGBColor(0xF4, 0xF6, 0xF8)


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def text_in(shape, value, size, bold=False, color=WHITE, align=PP_ALIGN.CENTER):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def box(slide, x, y, w, h, color, label, size=12):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(sh, color)
    sh.adjustments[0] = 0.1
    text_in(sh, label, size, True, WHITE)
    return sh


def label(slide, x, y, w, h, value, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    text_in(tb, value, size, bold, color, align)


def flow_row(slide, items, y, color=BLUE, w=1.45):
    x = 0.4
    for i, name in enumerate(items):
        box(slide, x, y, w, 0.55, color, name, 11)
        if i < len(items) - 1:
            label(slide, x + w, y + 0.05, 0.22, 0.45, "→", 16, True, NAVY, PP_ALIGN.CENTER)
        x += w + 0.2


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # Slide 1
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    fill(bg, NAVY)
    label(s, 0.7, 2.3, 12, 0.4, "CURRENT ARCHITECTURE", 16, True, RGBColor(0xC5, 0xD8, 0xE8))
    label(s, 0.7, 2.8, 12, 1.0, "Warehouse robotics control center", 36, True, WHITE)
    label(s, 0.7, 4.0, 12, 0.5, "Baseline 2.0.0  ·  from docs/03  ·  simple view", 16, False, RGBColor(0xD6, 0xE2, 0xEA))
    label(s, 0.7, 6.3, 12, 0.4, "No component has complete warehouse truth.", 18, True, AMBER)

    # Slide 2 — the only diagram
    s = prs.slides.add_slide(blank)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.85))
    fill(bar, NAVY)
    label(s, 0.5, 0.22, 12, 0.45, "How an order moves today", 24, True, WHITE)

    flow_row(s, ["ERP", "OMS", "WMS", "WES", "Fleet / WCS", "Robots", "Floor", "TMS"], 1.25, BLUE, 1.42)

    label(s, 0.5, 2.05, 12, 0.35, "Also connected", 13, True, SLATE)
    box(s, 0.5, 2.45, 2.4, 0.55, TEAL, "Vision / cameras", 12)
    box(s, 3.1, 2.45, 2.4, 0.55, TEAL, "Safety / IoT", 12)
    box(s, 5.7, 2.45, 3.3, 0.55, TEAL, "PLC / conveyor / ASRS", 12)
    box(s, 9.2, 2.45, 1.8, 0.55, TEAL, "Labor", 12)
    box(s, 11.2, 2.45, 1.6, 0.55, TEAL, "CMMS", 12)

    box(s, 0.5, 3.3, 12.3, 0.7, AMBER, "Shadow ops: email / Excel / radio  —  used when official systems lag", 14)

    label(s, 0.5, 4.3, 12, 2.4,
          "ERP books the qty.  OMS takes the order and cutoff.  WMS owns bins.  WES creates tasks.\n"
          "Fleet and WCS move robots and machines.  Floor is physical truth.  TMS books the truck.\n\n"
          "CMMS talks to fleet.  Vision and safety sit beside execution.  Humans still glue exceptions.",
          16, False, INK)

    # Slide 3 — one message
    s = prs.slides.add_slide(blank)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.85))
    fill(bar, NAVY)
    label(s, 0.5, 0.22, 12, 0.45, "What that means", 24, True, WHITE)

    box(s, 0.5, 1.25, 6.0, 2.4, BLUE, "Digital ≠ physical\n\nWMS, ERP, camera and email\ncan disagree on the same bin", 16)
    box(s, 6.8, 1.25, 6.0, 2.4, TEAL, "Available ≠ suitable\n\nFleet can say AVAILABLE while\ncert is expired or CMMS is open", 16)
    box(s, 0.5, 3.9, 12.3, 2.4, NAVY,
        "This repo observes. It does not drive robots.\nphysical_control = disabled",
        20)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
